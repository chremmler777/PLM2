"""MachineDB presses in the cost sheet (105): the MachineDB client, the sync
(plant mapping, idempotency, retired / scrapped / unmapped machines), the
per-machine rates in a draft (copied, frozen, diffed) and their precedence
over the class rate in costing."""
from datetime import date, datetime

import httpx
import pytest
from sqlalchemy import select

from app.models.change import ChangeAssessment, ChangeRequest
from app.models.cost_sheet import (
    CostSheetMachineClass, CostSheetMachineRate, CostSheetRate, CostSheetSamplingRate,
    CostSheetVersion,
)
from app.models.cost_sheet_machines import CostSheetMachine, CostSheetMachineItemRate
from app.models.entities import Plant, Project
from app.models.workflow import Department, UserDepartment
from app.services import cost_sheet_machines_service as msvc
from app.services import cost_sheet_service as cs
from app.services import machinedb_client as mdb

pytestmark = pytest.mark.asyncio

TOKEN = "s3cret-service-token"


def _row(id, name, plant, t, **kw):
    return {"id": id, "internal_name": name, "plant": plant,
            "clamping_force_t": None if t is None else f"{t:.2f}",
            "two_k_type": kw.get("two_k_type"), "manufacturer": kw.get("manufacturer", "Engel"),
            "model": kw.get("model"), "in_service_from": kw.get("in_service_from"),
            "planned_scrap_from": kw.get("planned_scrap_from"),
            "updated_at": "2026-09-01T00:00:00.000Z"}


# ---------------------------------------------------------------- client

@pytest.fixture
def configured(monkeypatch):
    monkeypatch.setenv("MACHINEDB_API_URL", "http://machinedb-backend:3001/v1/")
    monkeypatch.setenv("MACHINEDB_SERVICE_TOKEN", TOKEN)


async def test_client_reads_the_machine_list_with_the_bearer_token(configured):
    seen = {}

    def handler(request: httpx.Request):
        seen["url"] = str(request.url)
        seen["auth"] = request.headers.get("authorization")
        return httpx.Response(200, json=[
            _row(1, "P-80", "usa", 80, planned_scrap_from="2027-01-31"),
            _row(2, "P-350", "Mexico", 350.5, two_k_type="2k_turntable"),
            {"id": None, "internal_name": "broken"}])

    machines, bad = await mdb.list_machines(transport=httpx.MockTransport(handler))
    assert seen == {"url": "http://machinedb-backend:3001/v1/machines",
                    "auth": f"Bearer {TOKEN}"}
    assert [m.internal_name for m in machines] == ["P-80", "P-350"]
    assert machines[1].plant == "mexico" and machines[1].clamping_force_t == 350.5
    assert machines[1].two_k_type == "2k_turntable"
    assert machines[0].planned_scrap_from == date(2027, 1, 31)
    assert bad == ["row None"]


async def test_client_errors_never_carry_the_token(configured):
    import logging
    records = []

    class Keep(logging.Handler):
        def emit(self, record):
            records.append(record.getMessage())
    handler_ = Keep()
    logging.getLogger(mdb.__name__).addHandler(handler_)

    def refused(request):
        return httpx.Response(401, json={"error": "bad token " + TOKEN})

    def down(request):
        raise httpx.ConnectError(f"connection refused {request.headers['authorization']}")

    for handler, text in ((refused, "refused the service token"),
                          (down, "not reachable")):
        with pytest.raises(mdb.MachineDBUnavailable) as e:
            await mdb.list_machines(transport=httpx.MockTransport(handler))
        assert text in str(e.value) and TOKEN not in str(e.value)
    logging.getLogger(mdb.__name__).removeHandler(handler_)
    assert records and not any(TOKEN in r for r in records)


async def test_client_unconfigured(monkeypatch):
    monkeypatch.delenv("MACHINEDB_API_URL", raising=False)
    monkeypatch.setenv("MACHINEDB_SERVICE_TOKEN", TOKEN)
    assert not mdb.is_configured()
    assert mdb.config_status() == {"configured": False, "missing": ["MACHINEDB_API_URL"],
                                   "host": None}
    with pytest.raises(mdb.MachineDBUnavailable, match="not configured"):
        await mdb.list_machines()


async def test_config_status_shows_the_host_only(configured):
    assert mdb.config_status() == {"configured": True, "missing": [],
                                   "host": "machinedb-backend"}


# ---------------------------------------------------------------- world

@pytest.fixture
async def world(session_factory, seed):
    """Plants like the real ones: USA Toccoa (USD), Silao Mexico, Weissenburg
    (inactive), plus the seed's own plant. Machine classes by tonnage."""
    async with session_factory() as s:
        usa = Plant(organization_id=seed["org_id"], name="USA Toccoa", code="usa-toccoa",
                    location="Toccoa, GA, USA", is_active=True, currency="USD")
        mex = Plant(organization_id=seed["org_id"], name="Silao Mexico", code="SIL",
                    location="MX", is_active=True, currency="USD")
        wug = Plant(organization_id=seed["org_id"], name="Weissenburg", code="WUG",
                    location="DE", is_active=False, currency="EUR")
        small = CostSheetMachineClass(organization_id=seed["org_id"], name="<=200 t",
                                      tonnage_max=200, sort_order=0)
        big = CostSheetMachineClass(organization_id=seed["org_id"], name="200-450 t",
                                    tonnage_min=200, tonnage_max=450, sort_order=1)
        tool = Department(name="Tool Engineer", flow_type="action", is_active=True)
        sales = Department(name="Sales", flow_type="action", is_active=True)
        fin = Department(name="Finance", flow_type="info", is_active=True)
        s.add_all([usa, mex, wug, small, big, tool, sales, fin])
        await s.flush()
        return {**seed, "usa": usa.id, "mex": mex.id, "wug": wug.id, "small": small.id,
                "big": big.id, "tool": tool.id, "sales": sales.id, "fin": fin.id,
                "_commit": await s.commit()}


def _dtos(*rows):
    return [mdb.from_api(r) for r in rows]


FLEET = (_row(10, "P-80", "usa", 80), _row(11, "P-350", "usa", 350),
         _row(20, "M-300", "mexico", 300, two_k_type="2k_turntable"),
         _row(30, "W-200", "weissenburg", 200),
         _row(40, "S-500", "solingen", 500), _row(50, "B-120", "serbia", 120))


async def _sync(session_factory, org_id, rows, today=date(2026, 9, 26)):
    async with session_factory() as s:
        res = await msvc.sync_machines(s, org_id, None, machines=_dtos(*rows), today=today)
        await s.commit()
        return res


# ---------------------------------------------------------------- sync

async def test_sync_maps_plants_and_lists_unmapped(session_factory, world):
    res = await _sync(session_factory, world["org_id"], FLEET)
    assert res["total"] == 6 and res["new"] == 6 and res["changed"] == 0
    assert res["unmapped"] == 2
    assert res["report"]["unmapped"] == [
        {"machinedb_plant": "serbia", "count": 1, "machines": ["B-120"]},
        {"machinedb_plant": "solingen", "count": 1, "machines": ["S-500"]}]
    async with session_factory() as s:
        plants = {m.internal_name: m.plant_id for m in (await s.execute(
            select(CostSheetMachine))).scalars().all()}
        last = await msvc.last_sync(s, world["org_id"])
    assert plants == {"P-80": world["usa"], "P-350": world["usa"], "M-300": world["mex"],
                      "W-200": world["wug"], "S-500": None, "B-120": None}
    assert last["total"] == 6 and last["new"] == 6 and last["unmapped"] == 2


async def test_sync_is_idempotent_and_reports_changes(session_factory, world):
    await _sync(session_factory, world["org_id"], FLEET)
    again = await _sync(session_factory, world["org_id"], FLEET)
    assert (again["new"], again["changed"], again["retired"]) == (0, 0, 0)
    async with session_factory() as s:
        assert len((await s.execute(select(CostSheetMachine))).scalars().all()) == 6

    moved = list(FLEET)
    moved[1] = _row(11, "P-350", "usa", 400)                         # retonned
    moved[2] = _row(20, "M-300", "mexico", 300, two_k_type="2k_turntable",
                    planned_scrap_from="2026-06-30")                # scrapped
    del moved[5]                                                     # gone from MachineDB
    res = await _sync(session_factory, world["org_id"], moved)
    rep = res["report"]
    assert [(c["internal_name"], c["fields"]) for c in rep["changed"]] == [
        ("P-350", ["clamping_force_t"])]
    assert [c["internal_name"] for c in rep["scrapped"]] == ["M-300"]
    assert [c["internal_name"] for c in rep["retired"]] == ["B-120"]
    assert res["retired"] == 2
    async with session_factory() as s:
        rows = {m.internal_name: m for m in (await s.execute(
            select(CostSheetMachine))).scalars().all()}
    assert not rows["M-300"].active and rows["M-300"].retired_at is None
    assert not rows["B-120"].active and rows["B-120"].retired_at is not None
    assert rows["P-350"].clamping_force_t == 400

    # a machine back in MachineDB returns; a retired one is not reported twice
    back = await _sync(session_factory, world["org_id"], moved + [FLEET[5]])
    assert [c["internal_name"] for c in back["report"]["returned"]] == ["B-120"]
    assert back["report"]["retired"] == []


async def test_plant_mapping_setting_overrides_and_moves_machines(session_factory, world):
    await _sync(session_factory, world["org_id"], FLEET)
    async with session_factory() as s:
        eff = await msvc.set_plant_map(s, world["org_id"],
                                       {"solingen": world["wug"], "usa": None}, None)
        await s.commit()
    assert eff["solingen"] == {"plant_id": world["wug"], "source": "setting"}
    assert eff["usa"] == {"plant_id": None, "source": "setting"}
    assert eff["mexico"] == {"plant_id": world["mex"], "source": "default"}
    async with session_factory() as s:
        plants = {m.internal_name: m.plant_id for m in (await s.execute(
            select(CostSheetMachine))).scalars().all()}
    assert plants["S-500"] == world["wug"] and plants["P-80"] is None
    res = await _sync(session_factory, world["org_id"], FLEET)
    assert {u["machinedb_plant"] for u in res["report"]["unmapped"]} == {"usa", "serbia"}


async def test_default_mapping_never_guesses_between_two_plants(session_factory, world):
    async with session_factory() as s:
        s.add(Plant(organization_id=world["org_id"], name="Silao 2", code="SIL2",
                    location="Mexico", is_active=True))
        await s.commit()
        m = await msvc.plant_map(s, world["org_id"])
    assert m["mexico"] == {"plant_id": None, "source": "several plants match"}


# ---------------------------------------------------------------- API

async def test_sync_api_rights_and_unconfigured(client, admin_auth, eng_auth, world,
                                                session_factory, monkeypatch):
    monkeypatch.delenv("MACHINEDB_API_URL", raising=False)
    monkeypatch.delenv("MACHINEDB_SERVICE_TOKEN", raising=False)
    res = await client.post("/api/v1/cost-sheet/machines/sync", headers=eng_auth)
    assert res.status_code == 403
    res = await client.post("/api/v1/cost-sheet/machines/sync", headers=admin_auth)
    assert res.status_code == 503 and "not configured" in res.json()["detail"]
    listing = (await client.get("/api/v1/cost-sheet/machines", headers=eng_auth)).json()
    assert listing["machinedb"]["configured"] is False and listing["can_sync"] is False
    assert listing["last_sync"]["error"].startswith("MachineDB is not configured")
    # Sales may sync
    async with session_factory() as s:
        s.add(UserDepartment(user_id=world["engineer_id"], department_id=world["sales"]))
        await s.commit()

    async def fake_list():
        return _dtos(*FLEET[:2]), []
    monkeypatch.setattr(mdb, "list_machines", fake_list)
    res = await client.post("/api/v1/cost-sheet/machines/sync", headers=eng_auth)
    assert res.status_code == 200, res.text
    assert res.json()["new"] == 2
    listing = (await client.get("/api/v1/cost-sheet/machines", headers=eng_auth)).json()
    assert listing["can_sync"] is True
    assert [m["internal_name"] for m in listing["machines"]] == ["P-80", "P-350"]
    assert [m["machine_class"] for m in listing["machines"]] == ["<=200 t", "200-450 t"]
    assert listing["last_sync"]["new"] == 2 and "error" not in listing["last_sync"]


async def _published(session_factory, org_id, *, machines=(), items=(), sampling=(),
                     tool=None, version=1, valid_from=date(2026, 1, 1)):
    async with session_factory() as s:
        v = CostSheetVersion(organization_id=org_id, version=version, status="published",
                             valid_from=valid_from, published_at=datetime(2026, 1, 1))
        s.add(v)
        await s.flush()
        s.add(CostSheetRate(version_id=v.id, department_id=tool, hourly_rate=50,
                            currency="USD"))
        for m in machines:
            s.add(CostSheetMachineRate(version_id=v.id, **m))
        for mid, rate in items:
            s.add(CostSheetMachineItemRate(version_id=v.id, machine_id=mid,
                                           hourly_rate=rate, currency="USD"))
        for sm in sampling:
            s.add(CostSheetSamplingRate(version_id=v.id, **sm))
        await s.commit()
        return v.id


async def _machine_ids(session_factory):
    async with session_factory() as s:
        return {m.internal_name: m.id for m in (await s.execute(
            select(CostSheetMachine))).scalars().all()}


async def test_machine_rates_live_in_the_draft(client, admin_auth, world, session_factory):
    await _sync(session_factory, world["org_id"], FLEET)
    ids = await _machine_ids(session_factory)
    await _published(session_factory, world["org_id"], tool=world["tool"],
                     items=[(ids["P-80"], 70)])
    draft = (await client.post("/api/v1/cost-sheet/drafts", json={},
                               headers=admin_auth)).json()
    listing = (await client.get(f"/api/v1/cost-sheet/machines?version_id={draft['id']}",
                                headers=admin_auth)).json()
    rates = {m["internal_name"]: m["hourly_rate"] for m in listing["machines"]}
    assert rates["P-80"] == 70 and rates["P-350"] is None          # copied into the draft
    url = f"/api/v1/cost-sheet/machines/{ids['P-350']}/rate"
    res = await client.put(url, json={"version_id": draft["id"], "hourly_rate": 95.5},
                           headers=admin_auth)
    assert res.status_code == 200, res.text
    # the machine's plant is USA: its currency is the default
    assert res.json()["hourly_rate"] == 95.5 and res.json()["currency"] == "USD"
    res = await client.put(url, json={"version_id": draft["id"], "hourly_rate": 96},
                           headers=admin_auth)
    async with session_factory() as s:
        rows = (await s.execute(select(CostSheetMachineItemRate).where(
            CostSheetMachineItemRate.version_id == draft["id"]))).scalars().all()
    assert sorted(float(r.hourly_rate) for r in rows) == [70, 96]   # one row per machine
    d = (await client.get(f"/api/v1/cost-sheet/versions/{draft['id']}/diff",
                          headers=admin_auth)).json()
    assert d["machine_items"]["added"] == [{"machine_id": ids["P-350"], "hourly_rate": 96,
                                            "currency": "USD", "note": None,
                                            "entered_rate": None, "entered_currency": None}]
    # a draft whose only change is a machine rate publishes
    res = await client.post(f"/api/v1/cost-sheet/versions/{draft['id']}/publish",
                            json={"valid_from": "2027-01-01"}, headers=admin_auth)
    assert res.status_code == 200, res.text
    res = await client.put(url, json={"version_id": draft["id"], "hourly_rate": 99},
                           headers=admin_auth)
    assert res.status_code == 409                                   # frozen
    # removing needs a draft too; eng may not edit rates
    res = await client.put(url, json={"version_id": draft["id"], "hourly_rate": None},
                           headers=admin_auth)
    assert res.status_code == 409


async def test_machine_rate_refuses_bad_input(client, admin_auth, eng_auth, world,
                                              session_factory):
    await _sync(session_factory, world["org_id"], FLEET[:1])
    ids = await _machine_ids(session_factory)
    draft = (await client.post("/api/v1/cost-sheet/drafts", json={},
                               headers=admin_auth)).json()
    url = f"/api/v1/cost-sheet/machines/{ids['P-80']}/rate"
    assert (await client.put(url, json={"version_id": draft["id"], "hourly_rate": 5},
                             headers=eng_auth)).status_code == 403
    assert (await client.put(url, json={"version_id": draft["id"], "hourly_rate": -1},
                             headers=admin_auth)).status_code == 422
    assert (await client.put(url, json={"version_id": draft["id"], "hourly_rate": 5,
                                        "currency": "XXX"},
                             headers=admin_auth)).status_code == 422
    assert (await client.put("/api/v1/cost-sheet/machines/9999/rate",
                             json={"version_id": draft["id"], "hourly_rate": 5},
                             headers=admin_auth)).status_code == 404


# ---------------------------------------------------------------- costing

@pytest.fixture
async def change(session_factory, world):
    """A change in costing at USA Toccoa, Tool Engineer routed."""
    async with session_factory() as s:
        project = await s.get(Project, world["project_id"])
        project.plant_id = world["usa"]
        c = ChangeRequest(change_number="C-M-1", title="machines", reason="r",
                          change_type="physical_part", project_id=world["project_id"],
                          raised_by=world["admin_id"], lead_id=world["admin_id"],
                          status="costing")
        s.add(c)
        await s.flush()
        s.add(ChangeAssessment(change_id=c.id, department_id=world["tool"], stage_order=1,
                               verdict="feasible"))
        await s.commit()
        return c.id


async def _line(client, auth, change_id, world, **body):
    payload = {"department_id": world["tool"], "label": "Press", "kind": "machine_time",
               **body}
    res = await client.post(f"/api/v1/changes/{change_id}/costing/positions",
                            json=payload, headers=auth)
    assert res.status_code == 201, res.text
    return res.json()


async def test_machine_rate_beats_the_class_rate(client, admin_auth, world, change,
                                                 session_factory):
    await _sync(session_factory, world["org_id"], FLEET)
    ids = await _machine_ids(session_factory)
    await _published(session_factory, world["org_id"], tool=world["tool"],
                     machines=[dict(machine_class="200-450 t", machine_class_id=world["big"],
                                    plant_id=world["usa"], hourly_rate=85, currency="USD"),
                               dict(machine_class="200-450 t", machine_class_id=world["big"],
                                    hourly_rate=60, currency="USD")],
                     items=[(ids["P-350"], 120)])
    # the machine's own rate
    own = await _line(client, admin_auth, change, world, hours=2, machine_id=ids["P-350"])
    assert own["rate"] == 120 and own["line_value"] == 240
    assert own["rate_match"] == "machine" and own["machine_rate_own"]
    assert own["machine_name"] == "P-350" and own["machine_id"] == ids["P-350"]
    assert own["rate_label"] == "Cost sheet v1, Machine P-350, 120.00 USD/h"
    # no class on the line or the change: the machine's tonnage gives it
    assert own["machine_class_used_id"] == world["big"]
    assert not own["machine_class_from_change"]
    # a machine without its own rate falls back to the class rate of its plant
    m = await _line(client, admin_auth, change, world, hours=2, machine_id=ids["M-300"])
    assert m["rate"] == 60 and m["rate_match"] == "class"            # Silao has no class row
    # a line without a machine keeps the class pricing
    plain = await _line(client, admin_auth, change, world, hours=2,
                        machine_class_id=world["big"])
    assert plain["rate"] == 85 and plain["rate_match"] == "class+plant"
    assert plain["machine_id"] is None and not plain["machine_rate_own"]
    # choosing a machine later re-prices the line
    res = await client.put(f"/api/v1/changes/{change}/costing/positions/{plain['id']}",
                           json={"machine_id": ids["P-350"]}, headers=admin_auth)
    assert res.status_code == 200, res.text
    assert res.json()["rate"] == 120
    # a machine of another org is refused; labour lines drop the machine
    res = await client.post(f"/api/v1/changes/{change}/costing/positions",
                            json={"department_id": world["tool"], "label": "x",
                                  "kind": "machine_time", "machine_id": 99999},
                            headers=admin_auth)
    assert res.status_code == 400
    lab = await _line(client, admin_auth, change, world, kind="own_time", hours=1,
                      machine_id=ids["P-350"])
    assert lab["machine_id"] is None and lab["rate"] == 50


async def test_sampling_components_use_the_machine_rate(client, admin_auth, world, change,
                                                        session_factory):
    await _sync(session_factory, world["org_id"], FLEET)
    ids = await _machine_ids(session_factory)
    await _published(session_factory, world["org_id"], tool=world["tool"],
                     machines=[dict(machine_class="200-450 t", machine_class_id=world["big"],
                                    hourly_rate=80, currency="USD")],
                     items=[(ids["P-350"], 100)],
                     sampling=[dict(machine_class="200-450 t", machine_class_id=world["big"],
                                    mode="components", setup_hours=1, run_hours_default=2,
                                    handling_cost=10, currency="USD")])
    by_class = await _line(client, admin_auth, change, world, kind="sampling", trials=1,
                           machine_class_id=world["big"])
    assert by_class["rate"] == 250                                 # 3 h x 80 + 10
    by_machine = await _line(client, admin_auth, change, world, kind="sampling", trials=2,
                             machine_id=ids["P-350"])
    assert by_machine["rate"] == 310 and by_machine["line_value"] == 620  # 3 h x 100 + 10


async def test_sampling_with_machine_reports_a_foreign_currency():
    hit = cs.RateHit(rate=250, currency="USD", version_id=1, version=1, row_id=1,
                     match="class", breakdown={"mode": "components", "setup_hours": 1,
                                               "run_hours": 2, "labour_cost": 0,
                                               "handling_cost": 10, "missing": []})
    own = cs.RateHit(rate=90, currency="EUR", version_id=1, version=1, row_id=2,
                     match="machine")
    out = msvc.sampling_with_machine(hit, own)
    assert out.rate is None and out.breakdown["missing"] == ["machine_rate_currency"]
    flat = cs.RateHit(rate=99, currency="USD", version_id=1, version=1, row_id=1,
                      match="class", breakdown={"mode": "flat", "flat_price": 99})
    assert msvc.sampling_with_machine(flat, own) is flat



# ---------------------------------------------------------------- review fixes

async def test_plant_mapping_merges_and_a_key_can_be_dropped(client, admin_auth, world,
                                                             session_factory):
    await _sync(session_factory, world["org_id"], FLEET)
    async with session_factory() as s:
        await msvc.set_plant_map(s, world["org_id"], {"solingen": world["wug"]}, None)
        await s.commit()
    async with session_factory() as s:
        # a second mapping keeps the first
        eff = await msvc.set_plant_map(s, world["org_id"], {"serbia": world["mex"]}, None)
        await s.commit()
    assert eff["solingen"] == {"plant_id": world["wug"], "source": "setting"}
    assert eff["serbia"] == {"plant_id": world["mex"], "source": "setting"}
    # null unmaps on purpose; the others stay
    res = await client.put("/api/v1/cost-sheet/machines/plant-map",
                           json={"mapping": {"usa": None}}, headers=admin_auth)
    assert res.status_code == 200, res.text
    assert res.json()["usa"] == {"plant_id": None, "source": "setting"}
    assert res.json()["solingen"]["plant_id"] == world["wug"]
    # dropping a key brings the default back and moves the machines with it
    res = await client.delete("/api/v1/cost-sheet/machines/plant-map/usa", headers=admin_auth)
    assert res.status_code == 200, res.text
    assert res.json()["usa"] == {"plant_id": world["usa"], "source": "default"}
    assert res.json()["serbia"]["plant_id"] == world["mex"]
    assert (await client.delete("/api/v1/cost-sheet/machines/plant-map/usa",
                                headers=admin_auth)).status_code == 404
    plants = {m["internal_name"]: m["plant_id"] for m in (await client.get(
        "/api/v1/cost-sheet/machines", headers=admin_auth)).json()["machines"]}
    assert plants["P-80"] == world["usa"] and plants["B-120"] == world["mex"]


@pytest.mark.parametrize("answer, skipped, why", [
    ((), [], "listed no machines"),
    (FLEET[:2], ["row 1", "row 2", "row 3"], "unreadable rows"),
    (FLEET[:2], [], "would retire 4 of 6"),
])
async def test_a_broken_answer_retires_nothing_unless_forced(session_factory, world,
                                                             answer, skipped, why):
    await _sync(session_factory, world["org_id"], FLEET)
    async with session_factory() as s:
        with pytest.raises(cs.CostSheetError) as e:
            await msvc.sync_machines(s, world["org_id"], None, machines=_dtos(*answer),
                                     skipped=skipped, today=date(2026, 9, 26))
        await s.commit()
    assert e.value.status == 409 and why in e.value.message
    async with session_factory() as s:
        rows = (await s.execute(select(CostSheetMachine))).scalars().all()
        last = await msvc.last_sync(s, world["org_id"])
    assert len(rows) == 6 and all(r.retired_at is None for r in rows)
    assert last["guard"] is True and why in last["error"] and last["total"] == 6
    async with session_factory() as s:
        res = await msvc.sync_machines(s, world["org_id"], None, machines=_dtos(*answer),
                                       skipped=skipped, today=date(2026, 9, 26), force=True)
        await s.commit()
    assert res["retired"] == 6 - len(answer)


async def test_first_sync_and_small_changes_pass_the_guard(session_factory, world):
    assert (await _sync(session_factory, world["org_id"], ()))["total"] == 0
    await _sync(session_factory, world["org_id"], FLEET)
    res = await _sync(session_factory, world["org_id"], FLEET[:4])      # 2 of 6 go
    assert res["retired"] == 2


async def test_sync_api_force(client, admin_auth, world, session_factory, monkeypatch):
    await _sync(session_factory, world["org_id"], FLEET)

    async def empty():
        return [], []
    monkeypatch.setattr(mdb, "list_machines", empty)
    res = await client.post("/api/v1/cost-sheet/machines/sync", headers=admin_auth)
    assert res.status_code == 409 and "nothing was changed" in res.json()["detail"]
    listing = (await client.get("/api/v1/cost-sheet/machines", headers=admin_auth)).json()
    assert listing["last_sync"]["guard"] is True and len(listing["machines"]) == 6
    res = await client.post("/api/v1/cost-sheet/machines/sync?force=true", headers=admin_auth)
    assert res.status_code == 200 and res.json()["retired"] == 6


async def test_concurrent_writes_answer_409(client, admin_auth, world, session_factory,
                                            monkeypatch):
    from sqlalchemy.exc import IntegrityError
    await _sync(session_factory, world["org_id"], FLEET[:1])
    ids = await _machine_ids(session_factory)
    draft = (await client.post("/api/v1/cost-sheet/drafts", json={},
                               headers=admin_auth)).json()

    async def clash(*a, **k):
        raise IntegrityError("INSERT", {}, Exception("duplicate key"))
    monkeypatch.setattr(msvc, "set_machine_rate", clash)
    monkeypatch.setattr(msvc, "sync_machines", clash)
    res = await client.put(f"/api/v1/cost-sheet/machines/{ids['P-80']}/rate",
                           json={"version_id": draft["id"], "hourly_rate": 5},
                           headers=admin_auth)
    assert res.status_code == 409 and "same time" in res.json()["detail"]
    res = await client.post("/api/v1/cost-sheet/machines/sync", headers=admin_auth)
    assert res.status_code == 409 and "same time" in res.json()["detail"]


async def test_startup_sync_only_touches_configured_orgs(session_factory, world):
    from app.models.entities import Organization
    async with session_factory() as s:
        assert await msvc.startup_sync_orgs(s) == [world["org_id"]]     # the only org
        other = Organization(name="Other", code="other", is_active=True)
        s.add(other)
        await s.flush()
        s.add(Plant(organization_id=other.id, name="Elsewhere", code="ELS", is_active=True))
        await s.commit()
        assert await msvc.startup_sync_orgs(s) == []                    # two orgs: none
        await cs.set_setting(s, other.id, msvc.SETTING_MACHINEDB_ENABLED, "true")
        await s.commit()
        assert await msvc.startup_sync_orgs(s) == [other.id]
        await cs.set_setting(s, other.id, msvc.SETTING_MACHINEDB_ENABLED, "false")
        await cs.set_setting(s, world["org_id"], msvc.SETTING_MACHINEDB_ENABLED, "1")
        await s.commit()
        assert await msvc.startup_sync_orgs(s) == [world["org_id"]]


async def test_unmapped_machine_rate_needs_a_currency(client, admin_auth, world,
                                                      session_factory):
    await _sync(session_factory, world["org_id"], FLEET)
    ids = await _machine_ids(session_factory)
    draft = (await client.post("/api/v1/cost-sheet/drafts", json={},
                               headers=admin_auth)).json()
    url = f"/api/v1/cost-sheet/machines/{ids['S-500']}/rate"           # solingen: unmapped
    res = await client.put(url, json={"version_id": draft["id"], "hourly_rate": 80},
                           headers=admin_auth)
    assert res.status_code == 422 and "choose the currency" in res.json()["detail"]
    res = await client.put(url, json={"version_id": draft["id"], "hourly_rate": 80,
                                      "currency": "EUR", "note": "own press"},
                           headers=admin_auth)
    assert res.status_code == 200 and res.json()["currency"] == "EUR"
    assert res.json()["plant_currency"] is None
    # a later rate edit keeps the currency and the note
    res = await client.put(url, json={"version_id": draft["id"], "hourly_rate": 82},
                           headers=admin_auth)
    assert (res.json()["currency"], res.json()["note"]) == ("EUR", "own press")


async def test_machine_rate_in_another_currency_is_converted_or_missing(
        client, admin_auth, world, change, session_factory):
    """Costing at USA Toccoa (USD): P-350's own rate is in EUR, W-200 sits at
    Weissenburg whose class rate is EUR. Without a EUR/USD rate in the
    version both lines are unpriced; with one they are converted and say so."""
    await _sync(session_factory, world["org_id"], FLEET)
    ids = await _machine_ids(session_factory)
    async with session_factory() as s:
        v = CostSheetVersion(organization_id=world["org_id"], version=1, status="published",
                             valid_from=date(2026, 1, 1), published_at=datetime(2026, 1, 1))
        s.add(v)
        await s.flush()
        s.add_all([
            CostSheetRate(version_id=v.id, department_id=world["tool"], hourly_rate=50,
                          currency="USD"),
            CostSheetMachineRate(version_id=v.id, machine_class="<=200 t",
                                 machine_class_id=world["small"], plant_id=world["wug"],
                                 hourly_rate=40, currency="EUR"),
            CostSheetMachineItemRate(version_id=v.id, machine_id=ids["P-350"],
                                     hourly_rate=100, currency="EUR")])
        await s.commit()
        vid = v.id
    own = await _line(client, admin_auth, change, world, hours=2, machine_id=ids["P-350"])
    assert own["rate"] is None and own["rate_missing_reason"] == "machine_rate_currency"
    assert own["line_value"] is None
    cls = await _line(client, admin_auth, change, world, hours=1, machine_id=ids["W-200"])
    assert cls["rate"] is None and cls["rate_missing_reason"] == "machine_rate_currency"
    # the P&L flags them instead of dropping them silently
    from app.services.costing_rates import RateBook
    from app.services.pnl_service import PnlService
    async with session_factory() as s:
        c = await s.get(ChangeRequest, change)
        ctx = {"book": RateBook(s), "org": {c.id: world["org_id"]},
               "currency": {c.id: "USD"}, "plant": {c.id: world["usa"]}}
        _cost, no_rate = await PnlService.positions_cost(s, [c], ctx)
    assert no_rate == {change}
    # with the version's exchange rate: converted and labelled
    async with session_factory() as s:
        v = await s.get(CostSheetVersion, vid)
        v.fx_rates = {"EUR/USD": "1.10"}
        await s.commit()
    own = await _line(client, admin_auth, change, world, hours=2, machine_id=ids["P-350"])
    assert own["rate"] == 110 and own["rate_currency"] == "USD" and own["line_value"] == 220
    assert own["machine_rate_own"]
    assert "converted from 100.00 EUR" in own["rate_label"]
    cls = await _line(client, admin_auth, change, world, hours=1, machine_id=ids["W-200"])
    assert cls["rate"] == 44 and cls["rate_currency"] == "USD"


async def test_sampling_with_machine_converts_at_the_version_rate():
    hit = cs.RateHit(rate=250, currency="USD", version_id=1, version=1, row_id=1,
                     match="class", breakdown={"mode": "components", "setup_hours": 1,
                                               "run_hours": 2, "labour_cost": 0,
                                               "handling_cost": 10, "missing": []})
    own = cs.RateHit(rate=100, currency="EUR", version_id=1, version=1, row_id=2,
                     match="machine")
    v = CostSheetVersion(organization_id=1, version=1, fx_rates={"EUR/USD": "1.10"})
    out = msvc.sampling_with_machine(hit, own, v)
    assert out.rate == 340 and out.breakdown["machine_rate"] == 110      # 3 h x 110 + 10
    assert out.breakdown["machine_rate_fx"]["from_currency"] == "EUR"


async def test_machine_rate_typed_in_the_local_currency(client, admin_auth, world,
                                                        session_factory):
    """Silao quotes in USD and pays in MXN: a machine's own rate can be typed
    in MXN; USD is computed at the draft's exchange rate, the typed number is
    kept, and the export lists it."""
    async with session_factory() as s:
        mex = await s.get(Plant, world["mex"])
        mex.local_currency = "MXN"
        await s.commit()
    await _sync(session_factory, world["org_id"], FLEET)
    ids = await _machine_ids(session_factory)
    draft = (await client.post("/api/v1/cost-sheet/drafts", json={},
                               headers=admin_auth)).json()
    url = f"/api/v1/cost-sheet/machines/{ids['M-300']}/rate"
    body = {"version_id": draft["id"], "entered_rate": 1730, "entered_currency": "MXN"}
    res = await client.put(url, json=body, headers=admin_auth)
    assert res.status_code == 422 and "exchange rate" in res.json()["detail"]
    async with session_factory() as s:
        v = await s.get(CostSheetVersion, draft["id"])
        v.fx_rates = {"USD/MXN": "17.30"}
        await s.commit()
    res = await client.put(url, json=body, headers=admin_auth)
    assert res.status_code == 200, res.text
    m = res.json()
    assert (m["hourly_rate"], m["currency"]) == (100, "USD")
    assert (m["local_rate"], m["local_currency"], m["entered_in"]) == (1730, "MXN", "local")
    # a USA machine has no local column; MXN is not its currency
    usa = next(x for x in (await client.get(
        f"/api/v1/cost-sheet/machines?version_id={draft['id']}",
        headers=admin_auth)).json()["machines"] if x["internal_name"] == "P-80")
    assert usa["local_currency"] is None
    assert (await client.put(f"/api/v1/cost-sheet/machines/{ids['P-80']}/rate", json={
        **body}, headers=admin_auth)).status_code == 422
    # a new exchange rate re-prices it from the typed number
    async with session_factory() as s:
        v = await s.get(CostSheetVersion, draft["id"])
        await cs.set_fx_rate(s, v, "USD", "MXN", "17.00")
        await s.commit()
    m = next(x for x in (await client.get(
        f"/api/v1/cost-sheet/machines?version_id={draft['id']}",
        headers=admin_auth)).json()["machines"] if x["internal_name"] == "M-300")
    assert (m["hourly_rate"], m["local_rate"]) == (101.76, 1730)
    # typed in USD again: the MXN number is calculated, not kept
    res = await client.put(url, json={"version_id": draft["id"], "hourly_rate": 90},
                           headers=admin_auth)
    assert (res.json()["entered_in"], res.json()["local_rate"]) == ("quote", 1530)
    # the export lists the machine rates
    csv_text = (await client.get(
        f"/api/v1/cost-sheet/versions/{draft['id']}/export?format=csv&section=MachineRates",
        headers=admin_auth)).text
    assert "Machine;Plant;Tonnage t;Hourly rate;Currency" in csv_text
    assert "M-300;Silao Mexico;300,00;90,00;USD" in csv_text
    xlsx = await client.get(f"/api/v1/cost-sheet/versions/{draft['id']}/export?format=xlsx",
                            headers=admin_auth)
    import io
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(xlsx.content))
    assert "Machine rates" in wb.sheetnames
    assert [c.value for c in wb["Machine rates"][4]][:5] == ["M-300", "Silao Mexico", 300,
                                                            90, "USD"]
    # a rate in EUR at Silao: MXN typing asks to switch it to USD first
    res = await client.put(url, json={"version_id": draft["id"], "hourly_rate": 80,
                                      "currency": "EUR"}, headers=admin_auth)
    assert res.status_code == 200, res.text
    res = await client.put(url, json=body, headers=admin_auth)
    assert res.status_code == 422 and "switch it to USD first" in res.json()["detail"]
    res = await client.put(url, json={**body, "currency": "USD"}, headers=admin_auth)
    assert res.status_code == 200 and res.json()["currency"] == "USD"
    detail = (await client.get(f"/api/v1/cost-sheet/versions/{draft['id']}",
                               headers=admin_auth)).json()
    assert not [x for x in detail["currency_mismatch"] if x["section"] == "machine_items"]


async def test_sync_skips_an_out_of_range_clamping_force(session_factory, world):
    """A clamping force the column cannot hold is a skipped row in the
    report, never a failed sync; a known press is left as it was, not
    retired."""
    await _sync(session_factory, world["org_id"], FLEET)
    rows = [r for r in FLEET if r["id"] != 11] + [
        _row(11, "P-350", "usa", 1e9), _row(60, "X-neg", "usa", -5)]
    res = await _sync(session_factory, world["org_id"], rows)
    skipped = res["report"]["skipped"]
    assert any("P-350 (id 11)" in x and "out of range" in x for x in skipped)
    assert any("X-neg (id 60)" in x for x in skipped)
    assert not res["report"]["retired"] and not res["report"]["new"]
    async with session_factory() as s:
        p350 = (await s.execute(select(CostSheetMachine).where(
            CostSheetMachine.machinedb_id == 11))).scalar_one()
        assert p350.clamping_force_t == 350 and p350.retired_at is None
        assert (await s.execute(select(CostSheetMachine).where(
            CostSheetMachine.machinedb_id == 60))).scalar_one_or_none() is None
