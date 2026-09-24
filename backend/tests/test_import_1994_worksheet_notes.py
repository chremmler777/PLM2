"""The 1994 Excel import: plans comments, flags, material and Tier 1 numbers;
applies them once (a rerun adds nothing); never overwrites PLM cavities."""
import importlib.util
from pathlib import Path

import openpyxl
import pytest
from openpyxl.styles import PatternFill
from sqlalchemy import select

from app.models.field_note import FieldNote
from app.models.part import Part

pytestmark = pytest.mark.asyncio

SCRIPT = Path(__file__).resolve().parents[1] / "scripts" / "import_1994_worksheet_notes.py"
YELLOW, GREEN, SALMON = "FFFFFF00", "FFC6EFCE", "FFF8CBAD"
PREFIX = "From engineering Excel BOM 2026-09-23: "


def _load():
    spec = importlib.util.spec_from_file_location("import_1994_worksheet_notes", SCRIPT)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


def _xlsx(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "BOM"
    ws["A1"] = "Brose Seat Trim 1994"
    ws.append([])
    ws.append(["Project", "Tool no. (PLM)", "Cavities (PLM)", "Part no.", "BS part no."])
    row4 = {"A": "1994", "B": "199401", "C": "4", "D": "206.882.251", "E": "S00H4X-110",
            "F": "Handle, height adjustment LH", "I": "PA6-GF15 acc. VW 50125",
            "J": "Polykemi REZYcom PA6 RB122 F15", "K": "NOT OK: recycled; virgin grade to nominate",
            "L": "no", "M": "NM0", "N": "no", "T": "KF8", "U": "Stipple 4", "V": "2 ±0.3"}
    row5 = {"A": "1994", "B": "199403", "C": "2", "D": "206.887.233", "E": "S00H54-110",
            "F": "Isofix Cover", "I": "PA6-GF15 acc. VW 50125", "J": "Polykemi REZYcom PA6 RB122 F15",
            "K": "NOT OK: recycled", "L": "no", "M": "NM0", "N": "no (confirmed)",
            "R": "1 or 3 colors?", "S": "Confirmed 2026-09-23: no MIC, NM0 only",
            "T": "KF8", "U": "Stipple 2", "V": "2 ±0.3", "W": "RFQ has Stipple 2, drawing KF8"}
    fills = {4: {"F": GREEN, "K": SALMON}, 5: {"F": YELLOW, "C": YELLOW, "K": SALMON, "N": GREEN, "R": GREEN, "S": GREEN, "U": YELLOW, "W": YELLOW}}
    for r, data in ((4, row4), (5, row5)):
        for col, value in data.items():
            ws[f"{col}{r}"] = value
        for col, rgb in fills[r].items():
            ws[f"{col}{r}"].fill = PatternFill("solid", fgColor=rgb)
    path = tmp_path / "bom.xlsx"
    wb.save(path)
    return path


async def _parts(session_factory, seed):
    uid = seed["engineer_id"]
    async with session_factory() as s:
        def part(number, category, **kw):
            p = Part(project_id=seed["project_id"], part_number=number, name=number, part_type="internal_mfg",
                     item_category=category, created_by=uid, **kw)
            s.add(p)
            return p
        ids = {
            "lh": part("20-1994-001-0", "article", customer_part_number="206.882.251"),
            "iso": part("20-1994-005-0", "article", customer_part_number="206.887.233", tier1_part_number="S00H54-110"),
            "t01": part("199401", "tool", tool_cavities=2),
            "t03": part("199403", "tool", tool_cavities=4),
        }
        await s.commit()
        return {k: p.id for k, p in ids.items()}


async def test_plan_then_apply_once(session_factory, seed, tmp_path):
    mod = _load()
    ids = await _parts(session_factory, seed)
    rows = mod.read_rows(_xlsx(tmp_path))
    assert len(rows) == 2 and rows[1]["C"].flag == "open"

    async with session_factory() as s:
        actions, warnings = await mod.plan_actions(s, seed["project_id"], rows)
    assert warnings == []
    kinds = [(a.kind, a.part_id, a.field_key, a.flag) for a in actions]
    assert ("tier1", ids["lh"], None, None) in kinds
    assert not any(k == "tier1" and p == ids["iso"] for k, p, _, _ in kinds)  # already set in PLM
    assert ("material_new", ids["lh"], None, None) in kinds
    assert ("flag", ids["lh"], "part.material", "rejected") in kinds
    assert ("flag", ids["t01"], "tool.cavities", "open") in kinds
    assert ("flag", ids["t03"], "tool.cavities", "open") in kinds
    assert ("flag", ids["iso"], "paint.colour", "confirmed") in kinds
    assert ("flag", ids["iso"], "revision.level", "open") in kinds
    cav = [a.text for a in actions if a.kind == "comment" and a.field_key == "tool.cavities" and a.part_id == ids["t01"]]
    assert cav == [PREFIX + "4 cavities, PLM has 2. Check against RFQ 26 loop 37; PLM not changed."]
    assert all(a.text.startswith(PREFIX) for a in actions if a.kind == "comment")
    names = [(a.part_id, a.text) for a in actions if a.kind == "comment" and a.field_key == "part.name"]
    assert names == [(ids["iso"], PREFIX + "Designation on the drawing: Isofix Cover")]  # yellow only, not green

    async with session_factory() as s:
        written = await mod.apply_actions(s, actions, seed["engineer_id"])
        await s.commit()
    assert written == len(actions)
    async with session_factory() as s:
        lh = await s.get(Part, ids["lh"])
        t01 = await s.get(Part, ids["t01"])
        assert lh.tier1_part_number == "S00H4X-110"
        assert (lh.material_source, lh.material_new_text) == ("new", "PA6-GF15 acc. VW 50125")
        assert t01.tool_cavities == 2  # never overwritten
        notes = (await s.execute(select(FieldNote))).scalars().all()
        count = sum(len(n.comments) for n in notes)

    async with session_factory() as s:  # rerun: nothing new, the dry run already says so
        actions2, _ = await mod.plan_actions(s, seed["project_id"], rows)
        assert actions2 == []
        await mod.apply_actions(s, actions2, seed["engineer_id"])
        await s.commit()
        notes = (await s.execute(select(FieldNote))).scalars().all()
        assert sum(len(n.comments) for n in notes) == count
    assert not any(a.kind in ("tier1", "material_new") for a in actions2)


async def test_unknown_rows_are_warnings(session_factory, seed, tmp_path):
    mod = _load()
    rows = mod.read_rows(_xlsx(tmp_path))
    async with session_factory() as s:
        actions, warnings = await mod.plan_actions(s, seed["project_id"], rows)
    assert actions == []
    assert warnings == ["row 4: no article with OEM number 206.882.251", "row 5: no article with OEM number 206.887.233"]


async def test_comments_from_the_first_import_wording_are_not_planned_again(session_factory, seed, tmp_path):
    """plm_integ was imported before the prefix existed: a rerun must not add the same comment twice."""
    from app.services.field_note_service import FieldNoteService
    mod = _load()
    ids = await _parts(session_factory, seed)
    rows = mod.read_rows(_xlsx(tmp_path))
    async with session_factory() as s:
        t01 = await s.get(Part, ids["t01"])
        await FieldNoteService.add_comment(
            s, t01, "tool.cavities",
            "Excel BOM 2026-09-23 says 4 cavities, PLM has 2. Check against RFQ 26 loop 37; PLM not changed.",
            seed["engineer_id"])
        lh = await s.get(Part, ids["lh"])
        await FieldNoteService.add_comment(
            s, lh, "part.material",
            "Proposed resin: Polykemi REZYcom PA6 RB122 F15. Resin status: NOT OK: recycled; virgin grade to nominate",
            seed["engineer_id"])
        await s.commit()
    async with session_factory() as s:
        actions, _ = await mod.plan_actions(s, seed["project_id"], rows)
    planned = {(a.part_id, a.field_key) for a in actions if a.kind == "comment"}
    assert (ids["t01"], "tool.cavities") not in planned
    assert (ids["lh"], "part.material") not in planned


async def test_paint_questions_map_by_question_text():
    mod = _load()
    # "Painted or MIC?" and "Painted VM0 or chrome?" are about painted or not
    assert mod.QUESTION_FIELD["206.883.607"] == "paint.painted"
    assert mod.QUESTION_FIELD["206.881.793"] == "paint.painted"
    # "1 or 3 colors?" and "1 or 2 colors?" are about the colour
    assert mod.QUESTION_FIELD["206.887.233"] == "paint.colour"
    assert mod.QUESTION_FIELD["206.881.479"] == "paint.colour"


async def test_non_numeric_cavities_are_a_warning(session_factory, seed, tmp_path):
    mod = _load()
    ids = await _parts(session_factory, seed)
    rows = mod.read_rows(_xlsx(tmp_path))
    rows[0]["C"] = mod.Cell("tbd", None)
    async with session_factory() as s:
        actions, warnings = await mod.plan_actions(s, seed["project_id"], rows)
    assert warnings == ["row 4: cavities 'tbd' is not a number, skipped"]
    assert not any(a.part_id == ids["t01"] and a.field_key == "tool.cavities" for a in actions)


async def test_docstring_shows_user_as_required():
    doc = _load().__doc__
    assert "--user <id>" in doc and "[--user" not in doc


async def test_dry_run_disposes_the_engine(session_factory, seed, tmp_path, db_engine, monkeypatch, capsys):
    mod = _load()
    await _parts(session_factory, seed)
    from sqlalchemy.ext.asyncio import AsyncEngine
    disposed = []
    real_dispose = AsyncEngine.dispose

    async def dispose(self, *a, **k):
        if self is not db_engine:
            disposed.append(True)
        return await real_dispose(self, *a, **k)

    monkeypatch.setattr(AsyncEngine, "dispose", dispose)
    monkeypatch.setenv("DATABASE_URL", str(db_engine.url))
    monkeypatch.setattr("sys.argv", ["x", "--xlsx", str(_xlsx(tmp_path)), "--project", "proj",
                                     "--user", str(seed["engineer_id"])])
    await mod.main()
    assert disposed == [True]
    assert "DRY RUN" in capsys.readouterr().out
