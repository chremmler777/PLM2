"""Cost sheet (spec §15 / §15a): versions and validity chain, rights
(Sales, Finance or admin, acts-as aware), one rate per department per plant
(104) with empty rows seeded for every routable department, overheads,
machine and sampling prices, stale check, diff, export and the legacy
/changes/reference/rates shape."""
import io
from datetime import date, datetime, timedelta

import pytest
from sqlalchemy import select

from app.auth.acts_as import HEADER
from app.models.change_cost import DepartmentRate
from app.models.cost_sheet import (
    CostSheetMachineClass, CostSheetOverhead, CostSheetRate, CostSheetVersion,
    CostSheetMachineRate, CostSheetSamplingRate,
)
from app.models.entities import Plant
from app.models.workflow import Department, UserDepartment
from app.services import cost_sheet_service as svc

pytestmark = pytest.mark.asyncio
API = "/api/v1/cost-sheet"


@pytest.fixture
async def world(session_factory, seed):
    async with session_factory() as s:
        fin = Department(name="Finance", flow_type="info")
        tool = Department(name="Tooling-CS", flow_type="action")
        qa = Department(name="Quality-CS", flow_type="action")
        s.add_all([fin, tool, qa])
        await s.flush()
        p1 = (await s.execute(select(Plant))).scalars().first()
        p2 = Plant(organization_id=seed["org_id"], name="Second", code="p2", is_active=True)
        s.add(p2)
        await s.commit()
        return {**seed, "fin": fin.id, "tool": tool.id, "qa": qa.id,
                "p1": p1.id, "p2": p2.id}


async def _published(session_factory, org_id, *, version=1, valid_from=date(2026, 1, 1),
                     rates=(), overheads=(), machines=(), sampling=(),
                     published_at=None, note_override=None):
    async with session_factory() as s:
        v = CostSheetVersion(organization_id=org_id, version=version, status="published",
                             valid_from=valid_from, note=note_override,
                             published_at=published_at or datetime(2026, 1, 1))
        s.add(v)
        await s.flush()
        for r in rates:
            s.add(CostSheetRate(version_id=v.id, **r))
        for o in overheads:
            s.add(CostSheetOverhead(version_id=v.id, **o))
        for m in machines:
            s.add(CostSheetMachineRate(version_id=v.id, **m))
        for sm in sampling:
            s.add(CostSheetSamplingRate(version_id=v.id, **sm))
        await s.commit()
        return v.id


async def _load(session_factory, vid):
    async with session_factory() as s:
        return await s.get(CostSheetVersion, vid)


# ---------------------------------------------------------------- lookups

async def test_rate_is_the_department_row_at_the_plant(session_factory, world):
    t, q, p1, p2 = world["tool"], world["qa"], world["p1"], world["p2"]
    await _published(session_factory, world["org_id"], rates=[
        dict(department_id=t, hourly_rate=50),
        dict(department_id=t, plant_id=p1, hourly_rate=55),
        # an empty rate is no rate: Quality's all-plants row stands in at p1
        dict(department_id=q, plant_id=p1, hourly_rate=None),
        dict(department_id=q, hourly_rate=30),
        dict(department_id=world["fin"], plant_id=p1, hourly_rate=None),
    ])
    async with session_factory() as s:
        org = world["org_id"]
        # a position (older callers) is ignored: one rate per department
        hit = await svc.rate_for(s, org, t, "engineer", p1, date(2026, 5, 1))
        assert (hit.rate, hit.match) == (55, "department+plant")
        hit = await svc.rate_for(s, org, t, None, p2, date(2026, 5, 1))
        assert (hit.rate, hit.match) == (50, "department")
        assert hit.version == 1
        hit = await svc.rate_for(s, org, q, None, p1, date(2026, 5, 1))
        assert (hit.rate, hit.match) == (30, "department")
        # only an empty row: no rate, never 0
        assert await svc.rate_for(s, org, world["fin"], None, p1) is None
        assert await svc.effective_labour_rate(s, org, world["fin"], None, p1) is None
        # before the first version there is no rate at all
        assert await svc.rate_for(s, org, t, None, p1, date(2025, 12, 31)) is None


async def test_rate_uses_the_version_valid_on_the_date(session_factory, world):
    t = world["tool"]
    await _published(session_factory, world["org_id"], rates=[
        dict(department_id=t, hourly_rate=50)])
    await _published(session_factory, world["org_id"], version=2,
                     valid_from=date(2026, 7, 1),
                     rates=[dict(department_id=t, hourly_rate=60)])
    async with session_factory() as s:
        assert (await svc.rate_for(s, world["org_id"], t, on_date=date(2026, 6, 30))).rate == 50
        assert (await svc.rate_for(s, world["org_id"], t, on_date=date(2026, 7, 1))).rate == 60
        versions = await svc.list_versions(s, world["org_id"])
        valid = svc.validity(versions)
        assert valid[versions[0].id] == (date(2026, 1, 1), date(2026, 6, 30))
        assert valid[versions[1].id] == (date(2026, 7, 1), None)


async def test_effective_labour_rate_with_most_specific_overhead(session_factory, world):
    t, q, p1, p2 = world["tool"], world["qa"], world["p1"], world["p2"]
    await _published(session_factory, world["org_id"], rates=[
        dict(department_id=t, hourly_rate=100), dict(department_id=q, hourly_rate=100)],
        overheads=[
            dict(kind="percent", value=10),                           # org
            dict(plant_id=p2, kind="percent", value=20),              # plant
            dict(department_id=t, kind="per_hour", value=5),          # dept
            dict(department_id=t, plant_id=p2, kind="percent", value=50),  # dept+plant
        ])
    async with session_factory() as s:
        org = world["org_id"]
        assert (await svc.effective_labour_rate(s, org, t, None, p2)).rate == 150
        assert (await svc.effective_labour_rate(s, org, t, None, p1)).rate == 105
        assert (await svc.effective_labour_rate(s, org, q, None, p2)).rate == 120
        hit = await svc.effective_labour_rate(s, org, q, None, p1)
        assert (hit.rate, hit.base_rate, hit.overhead["kind"]) == (110, 100, "percent")


async def test_machine_rate_by_class_ref_and_plant(session_factory, world):
    p1, p2 = world["p1"], world["p2"]
    await _published(session_factory, world["org_id"], rates=[
        dict(department_id=world["tool"], hourly_rate=1)], machines=[
        dict(machine_class="200-450 t", hourly_rate=80),
        dict(machine_class="200-450 t", plant_id=p1, hourly_rate=90),
        dict(machine_class="200-450 t", machine_ref="Engel 350", hourly_rate=95),
    ])
    async with session_factory() as s:
        org = world["org_id"]
        assert (await svc.machine_rate_for(s, org, "200-450 t", p1)).rate == 90
        assert (await svc.machine_rate_for(s, org, "200-450 T", p2)).rate == 80
        assert (await svc.machine_rate_for(s, org, "200-450 t", p1,
                                           machine_ref="engel 350")).rate == 95
        assert (await svc.machine_rate_for(s, org, "200-450 t", p1,
                                           machine_ref="Other")).rate == 90
        assert await svc.machine_rate_for(s, org, ">800 t", p1) is None


async def test_sampling_flat_and_components(session_factory, world):
    t, p1 = world["tool"], world["p1"]
    await _published(session_factory, world["org_id"],
                     rates=[dict(department_id=t, position="Technician", hourly_rate=40)],
                     overheads=[dict(kind="percent", value=25)],
                     machines=[dict(machine_class="<=200 t", hourly_rate=60)],
                     sampling=[
                         dict(machine_class="<=200 t", mode="components", setup_hours=2,
                              run_hours_default=3, labour_hours=4, labour_department_id=t,
                              labour_position="Technician", handling_cost=100),
                         dict(machine_class=">800 t", mode="flat", flat_price=2500),
                     ])
    async with session_factory() as s:
        org = world["org_id"]
        hit = await svc.sampling_price_for(s, org, "<=200 t", p1)
        # (2+3) x 60 + 4 x 40 x 1.25 + 100
        assert hit.rate == 300 + 200 + 100
        assert hit.breakdown["missing"] == []
        hit = await svc.sampling_price_for(s, org, "<=200 t", p1, run_hours=5)
        assert hit.rate == 7 * 60 + 200 + 100
        assert (await svc.sampling_price_for(s, org, ">800 t", p1)).rate == 2500


async def test_class_for_tonnage():
    classes = [CostSheetMachineClass(name="<=200 t", tonnage_max=200),
               CostSheetMachineClass(name="200-450 t", tonnage_min=200, tonnage_max=450),
               CostSheetMachineClass(name=">800 t", tonnage_min=800)]
    assert svc.class_for_tonnage(classes, 150) == "<=200 t"
    assert svc.class_for_tonnage(classes, 200) == "<=200 t"
    assert svc.class_for_tonnage(classes, 350) == "200-450 t"
    assert svc.class_for_tonnage(classes, 600) is None
    assert svc.class_for_tonnage(classes, 1000) == ">800 t"


async def test_stale_status_follows_the_review_setting(session_factory, world):
    org = world["org_id"]
    async with session_factory() as s:
        st = await svc.stale_status(s, org)
        assert st["stale"] is True and st["reason"] == "no_published_version"
    await _published(session_factory, org, valid_from=date(2025, 1, 1),
                     published_at=datetime(2025, 1, 1),
                     rates=[dict(department_id=world["tool"], hourly_rate=1)])
    async with session_factory() as s:
        st = await svc.stale_status(s, org, today=date(2025, 12, 31))
        assert st["stale"] is False and st["due_on"] == "2026-01-01"
        assert (await svc.stale_status(s, org, today=date(2026, 1, 1)))["stale"] is True
        await svc.set_setting(s, org, "cost_sheet_review_months", "24")
        await s.commit()
        st = await svc.stale_status(s, org, today=date(2026, 6, 1))
        assert st["stale"] is False and st["review_months"] == 24


async def test_add_months_clamps():
    assert svc.add_months(date(2026, 1, 31), 1) == date(2026, 2, 28)
    assert svc.add_months(date(2026, 11, 15), 14) == date(2028, 1, 15)


# ---------------------------------------------------------------- API + rights

async def test_read_is_open_write_needs_finance(client, eng_auth, admin_auth, world,
                                                session_factory):
    res = await client.get(API, headers=eng_auth)
    assert res.status_code == 200
    body = res.json()
    assert body["can_edit"] is False and body["versions"] == []
    assert any(d["name"] == "Finance" for d in body["departments"])
    assert (await client.post(f"{API}/drafts", json={}, headers=eng_auth)).status_code == 403

    async with session_factory() as s:
        s.add(UserDepartment(user_id=world["engineer_id"], department_id=world["fin"]))
        await s.commit()
    assert (await client.get(API, headers=eng_auth)).json()["can_edit"] is True
    assert (await client.post(f"{API}/drafts", json={}, headers=eng_auth)).status_code == 201


async def test_sales_keeps_the_rates(client, eng_auth, admin_auth, world, session_factory):
    """Sales changes the rates: drafts, rows and publish; Finance and admins
    may too; anyone else reads."""
    res = await client.post(f"{API}/drafts", json={}, headers=eng_auth)
    assert res.status_code == 403
    assert res.json()["detail"] == "Only Sales, Finance or an admin may edit the cost sheet"
    async with session_factory() as s:
        sales = Department(name="Sales", flow_type="action")
        s.add(sales)
        await s.flush()
        s.add(UserDepartment(user_id=world["engineer_id"], department_id=sales.id))
        await s.commit()
        sales_id = sales.id
    assert (await client.get(API, headers=eng_auth)).json()["can_edit"] is True
    d = (await client.post(f"{API}/drafts", json={}, headers=eng_auth)).json()
    row = _row(d["rates"], world["tool"], None) or _row(d["rates"], world["tool"], world["p1"])
    res = await client.patch(f"{API}/versions/{d['id']}/rates/{row['id']}",
                             json={"hourly_rate": 42}, headers=eng_auth)
    assert res.status_code == 200
    res = await client.post(f"{API}/versions/{d['id']}/publish", json={
        "valid_from": "2026-01-01", "confirm_backdated": True}, headers=eng_auth)
    assert res.status_code == 200 and res.json()["status"] == "published"
    # the review task goes to Sales too
    body = (await client.get(f"{API}/review-task", headers=eng_auth)).json()
    assert body["is_editor"] is True
    as_sales = {**admin_auth, HEADER: str(sales_id)}
    assert (await client.get(API, headers=as_sales)).json()["can_edit"] is True


def _row(rows, department_id, plant_id):
    return next((r for r in rows if r["department_id"] == department_id
                 and r["plant_id"] == plant_id), None)


async def test_admin_acting_as_is_that_department_only(client, admin_auth, world):
    assert (await client.get(API, headers=admin_auth)).json()["can_edit"] is True
    as_tool = {**admin_auth, HEADER: str(world["tool"])}
    assert (await client.get(API, headers=as_tool)).json()["can_edit"] is False
    assert (await client.post(f"{API}/drafts", json={}, headers=as_tool)).status_code == 403
    as_fin = {**admin_auth, HEADER: str(world["fin"])}
    assert (await client.post(f"{API}/drafts", json={}, headers=as_fin)).status_code == 201


async def test_draft_edit_publish_freeze_and_chain(client, admin_auth, world):
    t, q, p1, p2 = world["tool"], world["qa"], world["p1"], world["p2"]
    d = (await client.post(f"{API}/drafts", json={}, headers=admin_auth)).json()
    assert d["version"] == 1 and d["status"] == "draft"
    # every routable department x plant starts with an EMPTY rate
    assert {(r["department_id"], r["plant_id"]) for r in d["rates"]} == {
        (dep, pl) for dep in (world["fin"], t, q) for pl in (p1, p2)}
    assert all(r["hourly_rate"] is None and r["effective_rate"] is None for r in d["rates"])
    # one draft at a time
    assert (await client.post(f"{API}/drafts", json={}, headers=admin_auth)).status_code == 409
    # a version without a single rate cannot be published
    res = await client.post(f"{API}/versions/{d['id']}/publish",
                            json={"valid_from": "2026-01-01"}, headers=admin_auth)
    assert res.status_code == 422

    res = await client.post(f"{API}/versions/{d['id']}/rates", json={
        "department_id": t, "hourly_rate": 50}, headers=admin_auth)
    assert res.status_code == 201, res.text
    # one row per department per plant: the seeded row is the one to fill
    res = await client.post(f"{API}/versions/{d['id']}/rates", json={
        "department_id": t, "plant_id": p1, "hourly_rate": 70}, headers=admin_auth)
    assert res.status_code == 409
    assert res.json()["detail"] == (
        "Tooling-CS already has a rate at Plant in this version. Edit that row instead.")
    res = await client.post(f"{API}/versions/{d['id']}/rates", json={
        "department_id": t, "hourly_rate": 1}, headers=admin_auth)
    assert res.status_code == 409 and "for all plants" in res.json()["detail"]
    # a position no longer makes a second row
    res = await client.post(f"{API}/versions/{d['id']}/rates", json={
        "department_id": t, "position": "Engineer", "plant_id": p1, "hourly_rate": 1},
        headers=admin_auth)
    assert res.status_code == 409
    seeded = _row(res_rates := (await client.get(f"{API}/versions/{d['id']}",
                                                 headers=admin_auth)).json()["rates"], t, p1)
    res = await client.patch(f"{API}/versions/{d['id']}/rates/{seeded['id']}", json={
        "hourly_rate": 70, "currency": "usd", "position": "Engineer"}, headers=admin_auth)
    row = _row(res.json()["rates"], t, p1)
    assert (row["hourly_rate"], row["currency"], row["position"]) == (70, "USD", None)
    # moving a row onto another department's plant is refused the same way
    other = _row(res_rates, q, p2)
    res = await client.patch(f"{API}/versions/{d['id']}/rates/{other['id']}",
                             json={"department_id": t, "plant_id": p1}, headers=admin_auth)
    assert res.status_code == 409
    default = _row(res.json() if res.status_code == 200 else
                   (await client.get(f"{API}/versions/{d['id']}", headers=admin_auth)
                    ).json()["rates"], t, None)
    # negative
    res = await client.patch(f"{API}/versions/{d['id']}/rates/{default['id']}",
                             json={"hourly_rate": -1}, headers=admin_auth)
    assert res.status_code == 422
    res = await client.patch(f"{API}/versions/{d['id']}/rates/{default['id']}",
                             json={"hourly_rate": 52.5}, headers=admin_auth)
    assert _row(res.json()["rates"], t, None)["hourly_rate"] == 52.5
    res = await client.post(f"{API}/versions/{d['id']}/overheads", json={
        "kind": "percent", "value": 10}, headers=admin_auth)
    assert _row(res.json()["rates"], t, None)["effective_rate"] == 57.75

    res = await client.post(f"{API}/versions/{d['id']}/publish", json={
        "valid_from": "2026-01-01", "note": "Budget 2026", "confirm_backdated": True},
        headers=admin_auth)
    assert res.status_code == 200
    assert res.json()["status"] == "published" and res.json()["valid_to"] is None
    # frozen: rows, publish and seeding
    res = await client.patch(f"{API}/versions/{d['id']}/rates/{default['id']}",
                             json={"hourly_rate": 1}, headers=admin_auth)
    assert res.status_code == 409
    assert (await client.post(f"{API}/versions/{d['id']}/missing-departments",
                              headers=admin_auth)).status_code == 409
    assert (await client.delete(f"{API}/versions/{d['id']}",
                                headers=admin_auth)).status_code == 409

    # v2 copies v1, must start after it
    d2 = (await client.post(f"{API}/drafts", json={}, headers=admin_auth)).json()
    assert d2["version"] == 2 and len(d2["rates"]) == 7 and len(d2["overheads"]) == 1
    res = await client.post(f"{API}/versions/{d2['id']}/publish",
                            json={"valid_from": "2026-01-01"}, headers=admin_auth)
    assert res.status_code == 422
    await client.patch(f"{API}/versions/{d2['id']}/rates/{_row(d2['rates'], t, None)['id']}",
                       json={"hourly_rate": 60}, headers=admin_auth)
    await client.delete(f"{API}/versions/{d2['id']}/rates/{_row(d2['rates'], t, p1)['id']}",
                        headers=admin_auth)
    await client.post(f"{API}/versions/{d2['id']}/rates", json={
        "department_id": q, "hourly_rate": 45}, headers=admin_auth)

    diff = (await client.get(f"{API}/versions/{d2['id']}/diff", headers=admin_auth)).json()
    assert diff["from_version"] == 1 and diff["to_version"] == 2
    assert diff["rates"]["changed"][0]["changes"]["hourly_rate"] == {"old": 52.5, "new": 60}
    assert diff["rates"]["changed"][0]["pct"] == 14.3
    assert [(r["department_id"], r["plant_id"]) for r in diff["rates"]["added"]] == [(q, None)]
    assert [(r["department_id"], r["plant_id"]) for r in diff["rates"]["removed"]] == [(t, p1)]
    assert diff["overheads"] == {"added": [], "removed": [], "changed": []}

    res = await client.post(f"{API}/versions/{d2['id']}/publish",
                            json={"valid_from": "2026-07-01", "confirm_backdated": True},
                            headers=admin_auth)
    assert res.status_code == 200
    # Tooling-CS has an all-plants row: publishing does not bring p1 back
    assert _row(res.json()["rates"], t, p1) is None
    ov = (await client.get(API, headers=admin_auth)).json()
    v1 = next(v for v in ov["versions"] if v["version"] == 1)
    assert v1["valid_to"] == "2026-06-30"
    assert ov["draft_version_id"] is None
    lookup = (await client.get(f"{API}/lookup/rate", params={
        "department_id": t, "on_date": "2026-03-01"}, headers=admin_auth)).json()
    assert lookup["rate"] == 57.75 and lookup["version"] == 1


async def test_add_missing_departments(client, admin_auth, eng_auth, world, session_factory):
    """The explicit action: empty rows for routable departments the draft
    does not cover, never a number; retired departments are left out; a
    department activated later is caught on publish too."""
    async with session_factory() as s:
        retired = Department(name="Retired-CS", flow_type="action", is_active=False)
        s.add(retired)
        await s.commit()
        retired_id = retired.id
    d = (await client.post(f"{API}/drafts", json={}, headers=admin_auth)).json()
    assert not any(r["department_id"] == retired_id for r in d["rates"])
    base = f"{API}/versions/{d['id']}"
    # Quality gets one rate for all plants instead of its two plant rows
    for pl in (world["p1"], world["p2"]):
        await client.delete(f"{base}/rates/{_row(d['rates'], world['qa'], pl)['id']}",
                            headers=admin_auth)
    await client.post(f"{base}/rates", json={"department_id": world["qa"], "hourly_rate": 40},
                      headers=admin_auth)
    await client.delete(f"{base}/rates/{_row(d['rates'], world['tool'], world['p2'])['id']}",
                        headers=admin_auth)
    assert (await client.post(f"{base}/missing-departments",
                              headers=eng_auth)).status_code == 403
    res = await client.post(f"{base}/missing-departments", headers=admin_auth)
    assert res.status_code == 200 and res.json()["added"] == 1
    added = _row(res.json()["rates"], world["tool"], world["p2"])
    assert added["hourly_rate"] is None and added["currency"] == "EUR"
    assert _row(res.json()["rates"], world["qa"], world["p1"]) is None
    assert (await client.post(f"{base}/missing-departments",
                              headers=admin_auth)).json()["added"] == 0
    # activated after the draft started: seeded when the version is published
    async with session_factory() as s:
        (await s.get(Department, retired_id)).is_active = True
        await s.commit()
    await client.patch(f"{base}/rates/{added['id']}", json={"hourly_rate": 10},
                       headers=admin_auth)
    res = await client.post(f"{base}/publish", json={
        "valid_from": "2026-01-01", "confirm_backdated": True}, headers=admin_auth)
    assert res.status_code == 200
    assert {r["plant_id"] for r in res.json()["rates"]
            if r["department_id"] == retired_id} == {world["p1"], world["p2"]}


async def test_delete_draft_and_versions_are_org_scoped(client, admin_auth, world,
                                                         session_factory):
    d = (await client.post(f"{API}/drafts", json={}, headers=admin_auth)).json()
    assert (await client.delete(f"{API}/versions/{d['id']}",
                                headers=admin_auth)).status_code == 204
    from app.models.entities import Organization
    async with session_factory() as s:
        other = Organization(name="Other", code="other", is_active=True)
        s.add(other)
        await s.flush()
        oid = other.id
        await s.commit()
    vid = await _published(session_factory, oid, rates=[
        dict(department_id=world["tool"], hourly_rate=1)])
    assert (await client.get(f"{API}/versions/{vid}", headers=admin_auth)).status_code == 404


async def test_row_validation(client, admin_auth, world):
    d = (await client.post(f"{API}/drafts", json={}, headers=admin_auth)).json()
    base = f"{API}/versions/{d['id']}"
    assert (await client.post(f"{base}/rates", json={"hourly_rate": 5},
                              headers=admin_auth)).status_code == 422
    assert (await client.post(f"{base}/rates", json={"department_id": world["tool"],
                                                     "hourly_rate": 5, "plant_id": 99999},
                              headers=admin_auth)).status_code == 422
    assert (await client.post(f"{base}/sampling", json={"machine_class": "x", "mode": "flat"},
                              headers=admin_auth)).status_code == 422
    assert (await client.post(f"{base}/sampling", json={
        "machine_class": "x", "mode": "components", "labour_hours": 2},
        headers=admin_auth)).status_code == 422
    # a row must name a class that exists
    assert (await client.post(f"{base}/machines", json={"machine_class": "<=200 t",
                                                        "hourly_rate": 60},
                              headers=admin_auth)).status_code == 422
    await client.post(f"{API}/machine-classes", json={"name": "<=200 t"}, headers=admin_auth)
    res = await client.post(f"{base}/machines", json={"machine_class": "<=200 t",
                                                      "hourly_rate": 60}, headers=admin_auth)
    assert res.status_code == 201
    res = await client.post(f"{base}/sampling", json={
        "machine_class": "<=200 t", "mode": "components", "setup_hours": 1,
        "run_hours_default": 1, "handling_cost": 10}, headers=admin_auth)
    assert res.json()["sampling_rates"][0]["computed_price"] == 130


async def test_machine_classes_and_settings(client, admin_auth, eng_auth, world):
    res = await client.post(f"{API}/machine-classes", json={
        "name": "<=200 t", "tonnage_max": 200}, headers=admin_auth)
    assert res.status_code == 201
    cid = res.json()["id"]
    assert (await client.post(f"{API}/machine-classes", json={"name": "<=200 t"},
                              headers=admin_auth)).status_code == 409
    assert (await client.post(f"{API}/machine-classes", json={"name": "x"},
                              headers=eng_auth)).status_code == 403
    await client.patch(f"{API}/machine-classes/{cid}", json={"sort_order": 3},
                       headers=admin_auth)
    assert (await client.get(f"{API}/machine-classes", headers=eng_auth)).json()[0][
        "sort_order"] == 3
    assert (await client.delete(f"{API}/machine-classes/{cid}",
                                headers=admin_auth)).status_code == 204
    assert (await client.get(f"{API}/machine-classes", headers=eng_auth)).json() == []

    assert (await client.get(f"{API}/settings", headers=eng_auth)).json() == {
        "review_months": 12}
    assert (await client.put(f"{API}/settings", json={"review_months": 6},
                             headers=eng_auth)).status_code == 403
    assert (await client.put(f"{API}/settings", json={"review_months": 6},
                             headers=admin_auth)).json() == {"review_months": 6}
    assert (await client.get(API, headers=eng_auth)).json()["stale"]["review_months"] == 6


async def test_export_csv_and_xlsx(client, admin_auth, world, session_factory):
    vid = await _published(session_factory, world["org_id"], rates=[
        dict(department_id=world["tool"], hourly_rate=70),
        dict(department_id=world["qa"], plant_id=world["p1"], hourly_rate=None)])
    res = await client.get(f"{API}/versions/{vid}/export", params={"format": "csv"},
                           headers=admin_auth)
    assert res.status_code == 200
    assert 'filename="cost-sheet-v1-rates.csv"' in res.headers["content-disposition"]
    text = res.content.decode("utf-8-sig")
    assert text.splitlines()[0].startswith("Department;Plant;Hourly rate")
    assert "Tooling-CS;All plants;70,00;70,00;EUR" in text
    assert "Quality-CS;Plant;;;EUR" in text            # an empty rate stays empty
    # the tab's old name still exports the rates
    old = await client.get(f"{API}/versions/{vid}/export",
                           params={"format": "csv", "section": "Positions"}, headers=admin_auth)
    assert old.content.decode("utf-8-sig") == text
    res = await client.get(f"{API}/versions/{vid}/export", headers=admin_auth)
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(res.content))
    assert wb.sheetnames[:4] == ["Rates", "Machines", "Sampling", "Overheads"]
    assert wb["Rates"]["B4"].value == "All plants"


async def test_reference_rates_reads_the_current_version(client, eng_auth, world,
                                                         session_factory):
    t, p1, p2 = world["tool"], world["p1"], world["p2"]
    async with session_factory() as s:
        s.add(DepartmentRate(department_id=t, plant_id=p1, hourly_rate=11.0,
                             min_factor=0.6, effective_from=date(2020, 1, 1)))
        await s.commit()
    legacy = (await client.get("/api/v1/changes/reference/rates", headers=eng_auth)).json()
    assert legacy == [{"department_id": t, "plant_id": p1, "hourly_rate": 11.0,
                       "min_factor": 0.6}]
    await _published(session_factory, world["org_id"], valid_from=date(2026, 1, 1), rates=[
        dict(department_id=t, plant_id=p1, hourly_rate=65, min_factor=0.6),
        dict(department_id=t, hourly_rate=50),
        dict(department_id=world["qa"], plant_id=p1, hourly_rate=None),   # empty: left out
    ])
    rows = (await client.get("/api/v1/changes/reference/rates", headers=eng_auth)).json()
    by_plant = {r["plant_id"]: r for r in rows}
    assert by_plant[p1] == {"department_id": t, "plant_id": p1, "hourly_rate": 65.0,
                            "min_factor": 0.6}
    assert by_plant[p2]["hourly_rate"] == 50.0 and by_plant[p2]["min_factor"] == 1.0
    assert len(rows) == 2


# ---------------------------------------------------------------- review fixes

async def _draft(client, auth):
    return (await client.post(f"{API}/drafts", json={}, headers=auth)).json()


async def test_nan_infinity_and_bounds_are_422(client, admin_auth, world):
    d = await _draft(client, admin_auth)
    base = f"{API}/versions/{d['id']}"
    hdr = {**admin_auth, "Content-Type": "application/json"}
    for raw in ('NaN', 'Infinity', '-Infinity', '1e9', '-1'):
        res = await client.post(f"{base}/rates", headers=hdr,
                                content=f'{{"department_id": {world["tool"]}, "hourly_rate": {raw}}}')
        assert res.status_code == 422, raw
    ok = await client.post(f"{base}/rates", json={"department_id": world["tool"],
                                                  "hourly_rate": 50}, headers=admin_auth)
    rid = ok.json()["rates"][0]["id"]
    res = await client.patch(f"{base}/rates/{rid}", headers=hdr, content='{"hourly_rate": NaN}')
    assert res.status_code == 422
    # the draft still loads: nothing broken was stored
    assert (await client.get(f"{base}", headers=admin_auth)).status_code == 200
    assert (await client.post(f"{base}/overheads", json={"kind": "percent", "value": 301},
                              headers=admin_auth)).status_code == 422
    assert (await client.post(f"{base}/overheads", json={"kind": "percent", "value": 300},
                              headers=admin_auth)).status_code == 201
    assert (await client.post(f"{base}/overheads", json={"kind": "per_hour", "value": -1,
                                                         "plant_id": world["p1"]},
                              headers=admin_auth)).status_code == 422
    await client.post(f"{API}/machine-classes", json={"name": "A"}, headers=admin_auth)
    res = await client.post(f"{base}/machines", json={"machine_class": "A", "hourly_rate": 1,
                                                      "tonnage_min": 500, "tonnage_max": 100},
                            headers=admin_auth)
    assert res.status_code == 422
    assert (await client.post(f"{base}/sampling", json={"machine_class": "A", "mode": "components",
                                                        "setup_hours": 1e7},
                              headers=admin_auth)).status_code == 422
    assert (await client.post(f"{API}/machine-classes", json={"name": "B", "tonnage_min": -5},
                              headers=admin_auth)).status_code == 422
    assert (await client.post(f"{API}/machine-classes", json={
        "name": "B", "tonnage_min": 9, "tonnage_max": 1}, headers=admin_auth)).status_code == 422


async def test_service_refuses_non_finite_numbers(session_factory, world):
    async with session_factory() as s:
        v = await svc.create_draft(s, world["org_id"], None)
        with pytest.raises(svc.CostSheetError) as e:
            await svc.add_row(s, v, "rates", {"department_id": world["tool"],
                                              "hourly_rate": float("nan")})
        assert e.value.status == 422


async def test_exports_neutralise_formulas(client, admin_auth, world, session_factory):
    vid = await _published(session_factory, world["org_id"], rates=[
        dict(department_id=world["tool"], hourly_rate=70, note="=HYPERLINK(\"x\")"),
        dict(department_id=world["tool"], plant_id=world["p1"], hourly_rate=70,
             note="+1+cmd|' /C calc'!A0"),
        dict(department_id=world["qa"], hourly_rate=12.5, note="@SUM(A1)"),
        dict(department_id=world["qa"], plant_id=world["p1"], hourly_rate=12.5, note="-2"),
    ], overheads=[dict(kind="percent", value=12.345)])
    csv_text = (await client.get(f"{API}/versions/{vid}/export",
                                 params={"format": "csv"}, headers=admin_auth)
                ).content.decode("utf-8-sig")
    assert "'=HYPERLINK" in csv_text and ";'+1+cmd" in csv_text and ";'@SUM" in csv_text
    assert ";'-2" in csv_text
    assert ";=" not in csv_text and ';"=' not in csv_text
    assert "70,00" in csv_text and "12,50" in csv_text           # decimal comma
    oh = (await client.get(f"{API}/versions/{vid}/export",
                           params={"format": "csv", "section": "Overheads"},
                           headers=admin_auth)).content.decode("utf-8-sig")
    assert "12,345" in oh                                        # full precision
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO((await client.get(f"{API}/versions/{vid}/export",
                                                    headers=admin_auth)).content))
    cells = [c.value for row in wb["Rates"].iter_rows(min_row=4) for c in row]
    assert "'=HYPERLINK(\"x\")" in cells and "'@SUM(A1)" in cells
    assert not any(isinstance(c, str) and c.startswith("=") for c in cells)


async def test_machine_class_rename_follows_rows(client, admin_auth, world):
    cls = (await client.post(f"{API}/machine-classes", json={"name": "small"},
                             headers=admin_auth)).json()
    await client.post(f"{API}/machine-classes", json={"name": "big"}, headers=admin_auth)
    d = await _draft(client, admin_auth)
    base = f"{API}/versions/{d['id']}"
    await client.post(f"{base}/rates", json={"department_id": world["tool"], "hourly_rate": 1},
                      headers=admin_auth)
    row = (await client.post(f"{base}/machines", json={"machine_class": "SMALL",
                                                       "hourly_rate": 40},
                             headers=admin_auth)).json()["machine_rates"][0]
    assert row["machine_class"] == "small" and row["machine_class_id"] == cls["id"]
    await client.post(f"{base}/publish", json={"valid_from": "2026-01-01",
                                               "confirm_backdated": True}, headers=admin_auth)
    assert (await client.patch(f"{API}/machine-classes/{cls['id']}", json={"name": "BIG"},
                               headers=admin_auth)).status_code == 409
    res = await client.patch(f"{API}/machine-classes/{cls['id']}", json={"name": "<=200 t"},
                             headers=admin_auth)
    assert res.status_code == 200 and res.json()["name"] == "<=200 t"
    got = (await client.get(f"{base}", headers=admin_auth)).json()
    assert got["machine_rates"][0]["machine_class"] == "<=200 t"
    hit = (await client.get(f"{API}/lookup/machine", params={"machine_class": "<=200 t",
                                                             "on_date": "2026-02-01"},
                            headers=admin_auth)).json()
    assert hit["rate"] == 40


async def test_currency_per_plant_and_no_mixing(client, admin_auth, world, session_factory):
    async with session_factory() as s:
        p2 = await s.get(Plant, world["p2"])
        p2.currency = "USD"
        await s.commit()
    ov = (await client.get(API, headers=admin_auth)).json()
    p2row = next(p for p in ov["plants"] if p["id"] == world["p2"])
    assert p2row["currency"] == "USD" and p2row["currency_confirmed"] is False
    d = await _draft(client, admin_auth)
    base = f"{API}/versions/{d['id']}"
    # the seeded row already carries the plant's currency
    assert _row(d["rates"], world["tool"], world["p2"])["currency"] == "USD"
    await client.delete(f"{base}/rates/{_row(d['rates'], world['tool'], world['p2'])['id']}",
                        headers=admin_auth)
    res = await client.post(f"{base}/rates", json={"department_id": world["tool"],
                                                   "plant_id": world["p2"], "hourly_rate": 30},
                            headers=admin_auth)
    assert _row(res.json()["rates"], world["tool"], world["p2"])["currency"] == "USD"
    assert (await client.post(f"{base}/rates", json={"department_id": world["qa"],
                                                     "hourly_rate": 1, "currency": "XYZ"},
                              headers=admin_auth)).status_code == 422
    # a per-hour overhead in EUR on a USD rate: no effective rate
    await client.post(f"{base}/overheads", json={"kind": "per_hour", "value": 5,
                                                 "currency": "EUR"}, headers=admin_auth)
    r = _row((await client.get(base, headers=admin_auth)).json()["rates"],
             world["tool"], world["p2"])
    assert r["effective_rate"] is None
    # Finance confirms the plant currency
    res = await client.put(f"{API}/plants/{world['p2']}/currency", json={"currency": "usd"},
                           headers=admin_auth)
    assert next(p for p in res.json() if p["id"] == world["p2"])["currency_confirmed"] is True
    assert (await client.put(f"{API}/plants/{world['p2']}/currency", json={"currency": "ABC"},
                             headers=admin_auth)).status_code == 422


async def test_components_sampling_incomplete_or_mixed_is_none(session_factory, world):
    t = world["tool"]
    await _published(session_factory, world["org_id"],
                     rates=[dict(department_id=t, hourly_rate=40, currency="USD")],
                     machines=[dict(machine_class="S", hourly_rate=60)],
                     sampling=[
                         dict(machine_class="S", mode="components", setup_hours=1,
                              labour_hours=2, labour_department_id=t),
                         dict(machine_class="M", mode="components", setup_hours=1),
                     ])
    async with session_factory() as s:
        hit = await svc.sampling_price_for(s, world["org_id"], "S")
        assert hit.rate is None and hit.breakdown["missing"] == ["labour_rate_currency"]
        assert hit.breakdown["complete"] is False
        hit = await svc.sampling_price_for(s, world["org_id"], "M")
        assert hit.rate is None and hit.breakdown["missing"] == ["machine_rate"]


async def test_backdated_publish_needs_confirmation(client, admin_auth, world):
    d = await _draft(client, admin_auth)
    await client.post(f"{API}/versions/{d['id']}/rates",
                      json={"department_id": world["tool"], "hourly_rate": 1}, headers=admin_auth)
    past = (date.today() - timedelta(days=3)).isoformat()
    res = await client.post(f"{API}/versions/{d['id']}/publish", json={"valid_from": past},
                            headers=admin_auth)
    assert res.status_code == 422 and "past" in res.json()["detail"]
    assert ("Changes created since then are priced with this version where a line "
            "has no rate yet, and hours booked since then use it too; lines already "
            "priced keep their rate.") in res.json()["detail"]
    res = await client.post(f"{API}/versions/{d['id']}/publish",
                            json={"valid_from": date.today().isoformat()}, headers=admin_auth)
    assert res.status_code == 200


async def test_publish_without_changes_is_409(client, admin_auth, world):
    d = await _draft(client, admin_auth)
    await client.post(f"{API}/versions/{d['id']}/rates",
                      json={"department_id": world["tool"], "hourly_rate": 1}, headers=admin_auth)
    await client.post(f"{API}/versions/{d['id']}/publish",
                      json={"valid_from": date.today().isoformat()}, headers=admin_auth)
    d2 = await _draft(client, admin_auth)
    res = await client.post(f"{API}/versions/{d2['id']}/publish",
                            json={"valid_from": (date.today() + timedelta(days=5)).isoformat()},
                            headers=admin_auth)
    assert res.status_code == 409 and res.json()["detail"] == "Nothing changed since version 1"


async def test_second_draft_blocked_by_the_database(client, admin_auth, world, monkeypatch):
    await _draft(client, admin_auth)

    async def no_draft(db, org_id):
        return None
    monkeypatch.setattr(svc, "open_draft", no_draft)       # simulate the race
    res = await client.post(f"{API}/drafts", json={}, headers=admin_auth)
    assert res.status_code == 409


async def test_migration_094_rebuilds_the_chain(db_engine, session_factory, world):
    import importlib.util
    import pathlib
    path = pathlib.Path(__file__).parents[1] / "alembic/versions/094_cost_sheet_fixes.py"
    import sys
    import types
    spec = importlib.util.spec_from_file_location("m094", path)
    m = importlib.util.module_from_spec(spec)
    # backend/alembic (the scripts dir) shadows the package; the helpers
    # under test never touch `op`.
    saved = sys.modules.get("alembic")
    sys.modules["alembic"] = types.SimpleNamespace(op=None)
    try:
        spec.loader.exec_module(m)
    finally:
        if saved is None:
            sys.modules.pop("alembic", None)
        else:
            sys.modules["alembic"] = saved
    assert m._is_usa("US", "Toccoa", "usa-toccoa") and m._is_usa(None, "USA", "x")
    assert not m._is_usa("DE", "Weissenburg", "WUG") and not m._is_usa("MX", "Silao", "SIL")
    t, p1, p2 = world["tool"], world["p1"], world["p2"]
    async with session_factory() as s:
        s.add_all([
            DepartmentRate(department_id=t, plant_id=p1, hourly_rate=50, min_factor=0.6,
                           effective_from=date(2026, 1, 1)),
            DepartmentRate(department_id=t, plant_id=p2, hourly_rate=20, min_factor=0.3,
                           effective_from=date(2026, 3, 1)),
            # same numbers again: no new version
            DepartmentRate(department_id=t, plant_id=p1, hourly_rate=50, min_factor=0.6,
                           effective_from=date(2026, 5, 1)),
            DepartmentRate(department_id=t, plant_id=p1, hourly_rate=55, min_factor=0.6,
                           effective_from=date(2026, 7, 1)),
        ])
        await s.commit()
    await _published(session_factory, world["org_id"], note_override=m.MIGRATED_NOTE,
                     rates=[dict(department_id=t, plant_id=p1, hourly_rate=55)])
    async with db_engine.begin() as conn:
        await conn.run_sync(lambda c: m._rebuild_migrated_chain(c, m.VERSIONS, {p2: "USD"}))
    async with session_factory() as s:
        vs = await svc.list_versions(s, world["org_id"])
        assert [(v.version, v.valid_from) for v in vs] == [
            (1, date(2026, 1, 1)), (2, date(2026, 3, 1)), (3, date(2026, 7, 1))]
        assert len(vs[0].rates) == 1 and len(vs[1].rates) == 2
        usd = next(r for r in vs[1].rates if r.plant_id == p2)
        assert usd.currency == "USD" and usd.hourly_rate == 20
        assert (await svc.rate_for(s, world["org_id"], t, None, p1, date(2026, 6, 30))).rate == 50
        assert (await svc.rate_for(s, world["org_id"], t, None, p1, date(2026, 7, 1))).rate == 55
