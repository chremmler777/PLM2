"""Builds the TOC-PLM-06 findings log (xlsx) from content_wi.py: the known
findings (section 9.1) and the Process Flow rules to confirm (section 8).
Run from build/: python3 make_findings_log.py"""
import re

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

import content_wi as W

HDR = PatternFill("solid", fgColor="D9D9D9")
THIN = Side(style="thin", color="BFBFBF")
WRAP = Alignment(wrap_text=True, vertical="top")
plain = lambda t: re.sub(r"\*\*", "", t or "")


def table_after(header_cell):
    """The rows of the content table whose header row contains header_cell."""
    for b in W.WI["blocks"]:
        if b[0] == "table" and b[2] and header_cell in b[2]:
            return b[3]
    raise KeyError(header_cell)


def sheet(ws, headers, widths, rows):
    ws.append(headers)
    for c, w in zip(ws[1], widths):
        c.font = Font(bold=True)
        c.fill = HDR
        c.alignment = WRAP
        ws.column_dimensions[c.column_letter].width = w
    for r in rows:
        ws.append(r)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = WRAP
            c.border = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def dropdown(ws, col, options, last):
    v = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=True)
    ws.add_data_validation(v)
    v.add(f"{col}2:{col}{last}")


def main():
    wb = Workbook()

    ws = wb.active
    ws.title = "Findings"
    known = []
    for no, text, where in table_after("Finding"):
        t = plain(text)
        status = ("Done" if "Fixed 10/01" in t else "Open")
        kind = "Process Flow" if no == "F-05" else ("Missing" if no == "F-04" else "Bug")
        known.append([no, "10/01/2026", "C. Demmler (desk check / walk)", where, "", "", kind,
                      "A" if no == "F-06" else "B", "", t, "", "", "", "", status])
    rows = known + [[f"S-{i:03d}"] + [None] * 14 for i in range(1, 201)]
    sheet(ws, ["ID", "Date", "Tester", "Simulation / step", "Change no.", "Acting as", "Type", "Severity",
               "What I did", "What happened (exact message)", "What I expected", "Screenshot file",
               "Decision (C. Demmler)", "Fix in version / LOP action", "Status"],
          [8, 11, 16, 13, 14, 16, 12, 9, 34, 38, 34, 18, 26, 18, 10], rows)
    last = len(rows) + 1
    dropdown(ws, "G", ["Bug", "Change", "Missing", "Text", "Process Flow"], last)
    dropdown(ws, "H", ["A", "B", "C"], last)
    dropdown(ws, "F", ["Myself (admin)", "Sales", "Project Manager", "Development", "Quality", "Scheduling",
                       "Tool Engineer", "APQP", "Manufacturing Engineer", "Packaging Engineer", "Finance",
                       "Other"], last)
    dropdown(ws, "O", ["Open", "Fix before go-live", "Fix after go-live", "No change", "Done"], last)

    ws2 = wb.create_sheet("Process Flow check")
    rules = table_after("Rule (Process Flow page and system, since 10/01/2026)")
    rows2 = [[no, plain(rule), where, None, None, None] for no, rule, where in rules]
    rows2 += [[f"M{i}", None, None, None, None, None] for i in range(len(rules) + 1, len(rules) + 13)]
    sheet(ws2, ["No.", "Rule (Process Flow page and system, since 10/01/2026)", "Check in step",
                "Seen in simulation (Yes / No / Other)", "Tester", "Comment (what you saw instead)"],
          [6, 70, 11, 16, 14, 44], rows2)
    dropdown(ws2, "D", ["Yes", "No", "Other"], len(rows2) + 1)

    ws3 = wb.create_sheet("Simulations run")
    sheet(ws3, ["Date", "Simulation (A / B / C1-C10)", "Change no.", "Person A", "Person B",
                "Completed to step", "Findings (IDs)", "Duration (h)"],
          [11, 18, 14, 16, 16, 14, 22, 11], [[None] * 8 for _ in range(40)])

    for w in wb.worksheets:
        w.sheet_view.zoomScale = 90
        w.page_setup.orientation = "landscape"
        w.page_setup.fitToWidth = 1
        w.page_setup.fitToHeight = 0
        w.sheet_properties.pageSetUpPr.fitToPage = True
    out = "../" + W.LOG
    wb.save(out)
    print(out)


if __name__ == "__main__":
    main()
