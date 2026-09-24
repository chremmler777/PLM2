"""xlsx export of the visible worksheet: typed numbers and dates, flag colours,
frozen identity columns, formula-like text kept as text."""
from datetime import date, datetime
from io import BytesIO

import openpyxl
import pytest

pytestmark = pytest.mark.asyncio

PAYLOAD = {
    "columns": [
        {"key": "part.part_number", "label": "KTX no.", "type": "text"},
        {"key": "part.customer_part_number", "label": "OEM no.", "type": "text"},
        {"key": "tool.cavities", "label": "Cavities", "type": "number"},
        {"key": "tool.cycle_time_s", "label": "Cycle time (s)", "type": "number"},
        {"key": "notes.summary", "label": "Notes", "type": "text"},
        {"key": "part.material_synced", "label": "Synced", "type": "date"},
    ],
    "rows": [
        {"cells": [{"value": "20-1994-001-0", "flag": None, "comments": 0},
                   {"value": "206.882.251", "flag": "confirmed", "comments": 0},
                   {"value": 2, "flag": "open", "comments": 1},
                   {"value": "55.5", "flag": None, "comments": 0},
                   {"value": "=1+1", "flag": "rejected", "comments": 2},
                   {"value": "2026-09-24T10:00:00", "flag": None, "comments": 0}]},
        {"cells": [{"value": "20-1994-002-0"}, {"value": None}, {"value": None},
                   {"value": "n/a"}, {"value": ""}, {"value": "not a date"}]},
    ],
    "frozen_columns": 2,
}


async def _export(client, auth, project_id, payload=PAYLOAD):
    return await client.post(f"/api/v1/projects/{project_id}/worksheet/export", json=payload, headers=auth)


async def test_export_writes_typed_cells_and_flag_colours(client, eng_auth, seed):
    r = await _export(client, eng_auth, seed["project_id"])
    assert r.status_code == 200, r.text
    assert r.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    today = date.today().isoformat()
    assert r.headers["content-disposition"] == f'attachment; filename="proj-worksheet-{today}.xlsx"'
    ws = openpyxl.load_workbook(BytesIO(r.content)).active
    assert [c.value for c in ws[1]] == ["KTX no.", "OEM no.", "Cavities", "Cycle time (s)", "Notes", "Synced"]
    assert ws["C2"].value == 2 and isinstance(ws["C2"].value, int)
    assert ws["D2"].value == 55.5
    assert isinstance(ws["F2"].value, datetime) and ws["F2"].value.date() == date(2026, 9, 24)
    assert ws["B2"].fill.fgColor.rgb == "FFC6EFCE"
    assert ws["C2"].fill.fgColor.rgb == "FFFFFF00"
    assert ws["E2"].fill.fgColor.rgb == "FFF8CBAD"
    assert ws["C2"].comment is not None and "1 comment" in ws["C2"].comment.text
    assert ws["D3"].value == "n/a"            # a number column keeps unparsable text
    assert ws["F3"].value == "not a date"
    assert ws["B3"].value is None
    assert ws.freeze_panes == "C2"
    assert ws.auto_filter.ref == "A1:F3"


async def test_export_keeps_formula_like_text_as_text(client, eng_auth, seed):
    r = await _export(client, eng_auth, seed["project_id"])
    ws = openpyxl.load_workbook(BytesIO(r.content)).active
    assert ws["E2"].value == "=1+1"
    assert ws["E2"].data_type == "s"


async def test_export_rejects_ragged_rows_and_foreign_projects(client, eng_auth, seed):
    bad = {**PAYLOAD, "rows": [{"cells": [{"value": "x"}]}]}
    assert (await _export(client, eng_auth, seed["project_id"], bad)).status_code == 422
    assert (await _export(client, eng_auth, 999999)).status_code == 404


def _one_cell(value, label="Notes", comments=0):
    return {"columns": [{"key": "notes.summary", "label": label, "type": "text"}],
            "rows": [{"cells": [{"value": value, "comments": comments}]}], "frozen_columns": 0}


async def test_export_caps_cell_text_label_and_comment_count(client, eng_auth, seed):
    pid = seed["project_id"]
    assert (await _export(client, eng_auth, pid, _one_cell("x" * 5000))).status_code == 200
    assert (await _export(client, eng_auth, pid, _one_cell("x" * 5001))).status_code == 422
    assert (await _export(client, eng_auth, pid, _one_cell("x", label="L" * 101))).status_code == 422
    assert (await _export(client, eng_auth, pid, _one_cell("x", comments=10_001))).status_code == 422


async def test_export_header_label_is_never_a_formula(client, eng_auth, seed):
    r = await _export(client, eng_auth, seed["project_id"], _one_cell("x", label='=HYPERLINK("http://evil","x")'))
    assert r.status_code == 200, r.text
    ws = openpyxl.load_workbook(BytesIO(r.content)).active
    assert ws["A1"].value == '=HYPERLINK("http://evil","x")'
    assert ws["A1"].data_type == "s"


async def test_export_strips_illegal_control_characters(client, eng_auth, seed):
    r = await _export(client, eng_auth, seed["project_id"], _one_cell("a\x01b\x1fc", label="La\x02bel"))
    assert r.status_code == 200, r.text
    ws = openpyxl.load_workbook(BytesIO(r.content)).active
    assert ws["A1"].value == "Label"
    assert ws["A2"].value == "abc"


async def test_build_xlsx_strips_control_characters_from_sheet_title():
    from app.services.worksheet_export import build_xlsx
    data = build_xlsx([{"key": "k", "label": "K", "type": "text"}], [[{"value": "v"}]], 0,
                      sheet_title="19\x0194 worksheet")
    wb = openpyxl.load_workbook(BytesIO(data))
    assert wb.active.title == "1994 worksheet"
