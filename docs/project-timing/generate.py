"""Project timing from standard blocks: G67 pilot.

Builds two files from one block library + one input table:
  G6x Timing Template.xlsx  - block library, team, inputs, formula-driven plan
  G67 Timeline (from template).xml - MS Project XML (MSPDI) with real links

Needs openpyxl, and for the XML mpxj + a Java runtime (JAVA_HOME).
Run: python generate.py
"""
import os
from datetime import date, datetime, timedelta

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

OUT = os.path.dirname(os.path.abspath(__file__))
XLSX = os.path.join(OUT, "G6x Timing Template.xlsx")
XML = os.path.join(OUT, "G67 Timeline (from template).xml")

# --- Team: role -> person (one row per project, filled once) -------------
TEAM = [
    ("Program Manager", "Steven Stocks", "Program Mgmt"),
    ("Plant Layout", "Stacy Sexton", "Plant Engineering"),
    ("Toolshop", "Brandon Belk", "Toolshop"),
    ("Gauge / Metrology", "Erica Gibson", "Quality"),
    ("CMM Lab", "Erika Harris", "Quality"),
    ("Trial Lead", "Christopher Holmes", "Process Engineering"),
    ("Process Engineer (WI)", "Scott Damewood", "Process Engineering"),
    ("Packaging", "Kenny Rice", "Logistics"),
    ("Quality Engineer", "Apurva Mhetre", "Quality"),
    ("Production Planning", "Alexa Earls", "Production"),
    ("Doc Control / ERP", "Todd Haynes", "Engineering"),
    ("Tool Engineering", "Matthias Rautzenberg", "Tool Engineering"),
    ("Purchasing", "Candy White", "Purchasing"),
    ("Press / Plant Eng", "Huascar Miranda", "Plant Engineering"),
]

# --- Block library ---------------------------------------------------------
# rule: ("input", col) start = input date
#       ("before", id)  finish the workday before id starts (scheduled backwards)
#       ("after", [ids]) start the workday after the latest finish
#       ("same", id)     start the same day as id
#       ("input_or_after", col, [ids])  input date if given, else after ids
#       ("max_input_after", col, [ids]) the later of input date and after ids
#       ("end", [ids])   milestone on the latest finish
CPK = "Weight; Thickness; every SPC characteristic on the control plan (from PDB)"
IPQ = ("0.1 Cover sheet; 0.2 Self-assessment; 1.3 Design/engineering approvals; "
       "1.4 IMDS; 2.1 Process flowchart; 2.2 PFMEA; 2.3 Control plan; "
       "3.1 Geometry/dimensions; 3.2 Material; 3.10 Reliability; 4.6 Tools; "
       "5.1 Legal compliance; 5.4 MSA; 5.5 Part history; 5.6 Load carriers")
BLOCKS = {
    "LAUNCH": ("New / transfer mold launch (one per mold)", [
        ("L01", "Assign location in plant layout", "Plant Layout", 15, ("before", "L03"), ""),
        ("L02", "Mold delivery", "Program Manager", 0, ("input", "delivery"), ""),
        ("L03", "Unpack / prep", "Toolshop", 2, ("after", ["L02"]), ""),
        ("L04", "Gauge validation", "Gauge / Metrology", 1, ("input_or_after", "gauge", ["L03"]), ""),
        ("L05", "TH1 trial", "Trial Lead", 1, ("max_input_after", "th1", ["L03", "L04"]), ""),
        ("L06", "Complete work instructions", "Process Engineer (WI)", 9, ("after", ["L05"]), "WI, first piece & safe launch WI, gauge WI"),
        ("L07", "TH2 trial", "Trial Lead", 1, ("after", ["L06"]), ""),
        ("L08", "Packaging trial / instruction", "Packaging", 1, ("same", "L07"), "PDS / packaging instruction"),
        ("L09", "6-pc dimensional (CMM)", "CMM Lab", 2, ("after", ["L07"]), "all cavities"),
        ("L10", "30-pc Cpk", "Quality Engineer", 4, ("after", ["L07"]), CPK),
        ("L11", "Mold ready for PPAP", "Program Manager", 0, ("end", ["L08", "L09", "L10"]), ""),
    ]),
    "MOLDCHG": ("Mold engineering change (belongs in the plm2 ECR plan)", [
        ("E01", "Bank build for the change", "Production Planning", 10, ("before", "E02"), "coverage until changed parts ship"),
        ("E02", "Ship tool to toolmaker", "Toolshop", 1, ("input", "tool_out"), ""),
        ("E03", "Tool modification", "Tool Engineering", 10, ("after", ["E02"]), "duration per quote"),
        ("E04", "Return tool to plant", "Toolshop", 1, ("after", ["E03"]), ""),
        ("E05", "Verify AI level engraved in mold", "Toolshop", 1, ("after", ["E04"]), ""),
        ("E06", "Change AI level in system", "Doc Control / ERP", 2, ("same", "E03"), ""),
        ("E07", "TH trial (changed tool)", "Trial Lead", 1, ("after", ["E05", "E06"]), ""),
        ("E08", "6-pc dimensional (CMM)", "CMM Lab", 2, ("after", ["E07"]), "all cavities"),
        ("E09", "30-pc Cpk", "Quality Engineer", 4, ("after", ["E07"]), CPK),
        ("E10", "Material disposition (old level)", "Production Planning", 1, ("same", "E07"), ""),
        ("E11", "Customer part approval / sampling", "Program Manager", 1, ("after", ["E08", "E09"]), "optional: vehicle trial"),
        ("E12", "Production start-up, new level", "Production Planning", 5, ("after", ["E11"]), ""),
        ("E13", "Changed parts delivered", "Program Manager", 0, ("end", ["E12"]), ""),
    ]),
    "PPAP": ("PPAP / ISIR package (one per part family)", [
        ("P01", "Initial capability (10-pc Cpk)", "Quality Engineer", 3, ("input", "trial"), CPK),
        ("P02", "ISIR / iPQ package", "Quality Engineer", 15, ("after", ["P01"]), IPQ),
        ("P03", "Validation testing (PV)", "Quality Engineer", 60, ("same", "P02"), "quote, select lab, ship samples, test"),
        ("P04", "Full PPAP approval", "Program Manager", 0, ("end", ["P02", "P03"]), ""),
    ]),
}

# --- G67 input: the only thing a PM types per mold -------------------------
# mold, press, parts, delivery, gauge validation, TH1 slot (press/trial plan)
# Values taken from G6x Project Timeline (2026-09-30).mpp.
D = date
MOLDS = [
    ("3355", "900-5", "5A65DF8-03", D(2026, 8, 10), D(2026, 10, 2), D(2026, 10, 7)),
    ("3363", "1300-3", "5A65E17-03, 5A65E18-03", D(2026, 9, 3), D(2026, 10, 5), D(2026, 10, 8)),
    ("3370", "1600-2", "5B533A6-02", D(2026, 9, 3), D(2026, 10, 6), D(2026, 10, 9)),
    ("3377", "1300-2", "9635313-02, 9635314-02", D(2026, 8, 10), D(2026, 10, 7), D(2026, 10, 13)),
    ("3376", "1300-3", "9635311-02, 9635312-02", D(2026, 8, 10), D(2026, 10, 14), None),
    ("3362", "1600-2", "5A65E00-03, 5A65E00-04", D(2026, 9, 3), D(2026, 10, 15), D(2026, 10, 15)),
    ("3343", "1600-2", "5A87904-02, 5A65C28-02", D(2026, 9, 3), D(2026, 10, 15), D(2026, 10, 16)),
    ("3347", "1300-3", "5B6CA70-04", D(2026, 8, 10), D(2026, 10, 16), D(2026, 10, 20)),
    ("3352", "550", "5B6CA81-02, 5B6CA82-02", D(2026, 9, 3), D(2026, 10, 19), D(2026, 10, 22)),
]
# PM's hand-typed (start, finish) per mold and step, for the comparison columns
PM = {
    "3355": {"L03": ("08-11", "08-12"), "L05": ("10-07", "10-07"), "L06": ("10-08", "10-20"), "L07": ("10-21", "10-21"), "L09": ("10-22", "10-23"), "L10": ("10-22", "10-27")},
    "3363": {"L03": ("09-04", "09-07"), "L05": ("10-08", "10-08"), "L06": ("10-09", "10-21"), "L07": ("10-22", "10-22"), "L09": ("10-23", "10-26"), "L10": ("10-23", "10-28")},
    "3370": {"L03": ("09-04", "09-07"), "L05": ("10-09", "10-09"), "L06": ("10-12", "10-22"), "L07": ("10-23", "10-23"), "L09": ("10-26", "10-27"), "L10": ("10-26", "10-29")},
    "3377": {"L03": ("08-17", "08-18"), "L05": ("10-13", "10-13"), "L06": ("10-14", "10-26"), "L07": ("10-27", "10-27"), "L09": ("10-28", "10-29"), "L10": ("10-28", "11-02")},
    "3376": {"L03": ("08-13", "08-14"), "L06": ("10-05", "10-13"), "L07": ("10-14", "10-14"), "L09": ("10-15", "10-16"), "L10": ("10-15", "10-20")},
    "3362": {"L03": ("08-31", "09-01"), "L05": ("10-15", "10-15"), "L06": ("10-16", "10-28"), "L07": ("10-29", "10-29"), "L09": ("10-30", "11-05"), "L10": ("10-30", "11-04")},
    "3343": {"L03": ("08-31", "09-01"), "L05": ("10-16", "10-16"), "L06": ("10-19", "10-29"), "L07": ("10-30", "10-30"), "L09": ("11-02", "11-03"), "L10": ("11-02", "11-05")},
    "3347": {"L03": ("08-19", "08-20"), "L05": ("10-20", "10-20"), "L06": ("10-21", "11-02"), "L07": ("11-03", "11-03"), "L09": ("11-04", "11-05"), "L10": ("11-04", "11-09")},
    "3352": {"L03": ("09-02", "09-03"), "L05": ("10-22", "10-22"), "L06": ("10-23", "11-04"), "L07": ("11-05", "11-05"), "L09": ("11-06", "11-09"), "L10": ("11-06", "11-11")},
}
FINDINGS = [
    ("Links", "0 of 239 tasks are linked, no constraints, no baseline", "Every date is typed by hand; a slip moves nothing and cannot be measured"),
    ("Logic", "Molds 3362, 3343, 3352: unpack (8/31-9/3) planned before delivery (9/3)", "A link delivery -> unpack makes this impossible"),
    ("Logic", "Mold 3376: no TH1, TH2 on the same day as gauge validation", "Template always plans TH1 after gauge"),
    ("Logic", "Mold 3362: TH1 on the same day as gauge validation", "Template starts TH1 the workday after gauge"),
    ("Naming", "Mold 3377: press 1300-2 in some lines, 1300-3 in others", "Press is one input field per mold"),
    ("Naming", "Task 5.9.11 (3352) called 'TH2' but is the 30-pc Cpk", "Task names come from the block"),
    ("Resources", "Same person 3x: holmes / Chrisopher.holmes / christopher.holmes; Rautzenberg 3x", "Roles in the block, one role -> person map per project"),
    ("Resources", "A resource named 'Mold 3345: Verify correct AI Level is engraved in the mold.'", "Same"),
    ("Granularity", "Cpk split into 8 rows per characteristic (24 rows total)", "One Cpk task, characteristics as checklist from the control plan / PDB"),
    ("Granularity", "15 iPQ element rows with identical dates and owner", "One ISIR/iPQ task with the element checklist"),
    ("Scope", "Mold 3360/3349/3369 changes (~55 rows) planned in the project file", "These are ECRs: plan them in the plm2 change plan, show only milestones here"),
]

ROLE_COL = {"delivery": "D", "gauge": "E", "th1": "F"}
HOL = "Holidays!$A$2:$A$60"


# --- Python mirror of the Excel rules (dates for the MS Project XML) -------
def workday(d, n):
    step = 1 if n >= 0 else -1
    while n:
        d += timedelta(days=step)
        if d.weekday() < 5:
            n -= step
    return d


def schedule(steps, inputs):
    s, f = {}, {}
    pending = list(steps)
    while pending:
        for st in list(pending):
            sid, _, _, dur, rule, _ = st
            kind = rule[0]
            deps = {"after": rule[1] if kind == "after" else [], "same": [rule[1]] if kind == "same" else [],
                    "before": [rule[1]] if kind == "before" else [], "end": rule[1] if kind == "end" else [],
                    "input_or_after": rule[2] if kind == "input_or_after" else [],
                    "max_input_after": rule[2] if kind == "max_input_after" else []}[kind] if kind != "input" else []
            if any(d not in f for d in deps):
                continue
            if kind == "input":
                start = inputs[rule[1]]
            elif kind == "before":
                start = workday(s[rule[1]], -dur)
            elif kind == "same":
                start = s[rule[1]]
            elif kind == "end":
                start = max(f[d] for d in deps)
            elif kind == "input_or_after" and inputs.get(rule[1]):
                start = inputs[rule[1]]
            else:
                start = max(workday(f[d], 1) for d in deps)
                if kind == "max_input_after" and inputs.get(rule[1]):
                    start = max(start, inputs[rule[1]])
            s[sid], f[sid] = start, (start if dur == 0 else workday(start, dur - 1))
            pending.remove(st)
    return s, f


# --- Excel -----------------------------------------------------------------
FONT = "Arial"
HDR = PatternFill("solid", start_color="1F3864")
BAND = PatternFill("solid", start_color="D9E1F2")
INPUT = Font(name=FONT, color="0000FF")
LINK = Font(name=FONT, color="008000")
THIN = Border(bottom=Side(style="thin", color="BFBFBF"))


def header(ws, row, cols, widths):
    for i, (c, w) in enumerate(zip(cols, widths), 1):
        cell = ws.cell(row=row, column=i, value=c)
        cell.font = Font(name=FONT, bold=True, color="FFFFFF")
        cell.fill = HDR
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = ws.cell(row=row + 1, column=1)


def style_rows(ws, first, last, ncols):
    for r in range(first, last + 1):
        for c in range(1, ncols + 1):
            cell = ws.cell(row=r, column=c)
            if cell.font.color is None or cell.font.color.rgb in (None, "FF000000"):
                cell.font = Font(name=FONT, bold=cell.font.bold)
            else:
                cell.font = Font(name=FONT, bold=cell.font.bold, color=cell.font.color)
            cell.border = THIN
            cell.alignment = Alignment(vertical="top", wrap_text=c in (3, 8) and ws.title == "Blocks")


def build_xlsx():
    wb = Workbook()
    rd = wb.active
    rd.title = "README"
    lines = [
        ("G6x Timing Template: standard blocks instead of copy-paste", True),
        ("", False),
        ("How it works", True),
        ("1. Blocks: the standard task sequences (mold launch, mold change, PPAP), with roles, durations and logic. Owned centrally, versioned.", False),
        ("2. Team: one role -> person map per project. Blocks never name people.", False),
        ("3. Inputs: per mold only press, part numbers, delivery date, gauge validation date and the TH1 trial slot (blue cells).", False),
        ("4. Plan: generated from the three sheets above. All dates are formulas: change an input or a duration and everything after it moves.", False),
        ("5. Project-specific tasks (customer trials, downtime, payment setup...) are planned by hand on top; they are not in the blocks.", False),
        ("", False),
        ("Rules", True),
        ("- Workdays Mon-Fri; plant holidays go on the Holidays sheet.", False),
        ("- Cpk and iPQ elements are checklists on one task, not one row per characteristic / element.", False),
        ("- Mold engineering changes are ECRs: plan them in the plm2 change plan; the project shows only their milestones.", False),
        ("- 'PM plan' columns hold the dates from G6x Project Timeline (2026-09-30).mpp for comparison; 'Delta' = template finish minus PM finish (calendar days).", False),
        ("", False),
        ("Colors: blue = input, black = formula, green = pulled from another sheet.", False),
        ("Source: G6x Project Timeline (2026-09-30).mpp, PM 1748, read 2026-10-02.", False),
    ]
    for i, (t, b) in enumerate(lines, 1):
        rd.cell(row=i, column=1, value=t).font = Font(name=FONT, bold=b, size=14 if i == 1 else 10)
    rd.column_dimensions["A"].width = 130

    tm = wb.create_sheet("Team")
    header(tm, 1, ["Role", "Person", "Department"], [26, 24, 22])
    for i, row in enumerate(TEAM, 2):
        tm.cell(row=i, column=1, value=row[0])
        for c in (2, 3):
            tm.cell(row=i, column=c, value=row[c - 1]).font = INPUT
    style_rows(tm, 2, len(TEAM) + 1, 3)

    bl = wb.create_sheet("Blocks")
    header(bl, 1, ["Block", "Step ID", "Step", "Role", "Duration (wd)", "Depends on", "Rule", "Checklist / deliverable"],
           [11, 8, 34, 22, 12, 16, 40, 60])
    rule_txt = {
        "input": lambda r: f"Start = input '{r[1]}'",
        "before": lambda r: f"Finish the workday before {r[1]} starts",
        "after": lambda r: "Start the workday after " + " and ".join(r[1]) + " finish",
        "same": lambda r: f"Start the same day as {r[1]}",
        "input_or_after": lambda r: f"Input '{r[1]}' if given, else after " + " and ".join(r[2]),
        "max_input_after": lambda r: f"Later of input '{r[1]}' and after " + " and ".join(r[2]),
        "end": lambda r: "Milestone when " + ", ".join(r[1]) + " are done",
    }
    brow, r = {}, 2
    for key, (title, steps) in BLOCKS.items():
        c = bl.cell(row=r, column=1, value=f"{key}: {title}")
        c.font = Font(name=FONT, bold=True)
        for col in range(1, 9):
            bl.cell(row=r, column=col).fill = BAND
        r += 1
        for sid, name, role, dur, rule, chk in steps:
            deps = rule[1] if rule[0] in ("after", "end") else rule[2] if rule[0] in ("input_or_after", "max_input_after") else [rule[1]] if rule[0] in ("before", "same") else []
            vals = [key, sid, name, role, dur, ", ".join(deps), rule_txt[rule[0]](rule), chk]
            for col, v in enumerate(vals, 1):
                cell = bl.cell(row=r, column=col, value=v)
                if col == 5:
                    cell.font = INPUT
            brow[sid] = r
            r += 1
    style_rows(bl, 2, r - 1, 8)
    dv = DataValidation(type="list", formula1=f"=Team!$A$2:$A${len(TEAM) + 1}", allow_blank=False)
    bl.add_data_validation(dv)
    dv.add(f"D2:D{r - 1}")

    inp = wb.create_sheet("Inputs")
    header(inp, 1, ["Mold", "Press", "Part numbers", "Delivery", "Gauge validation", "TH1 trial slot"],
           [10, 10, 28, 14, 16, 15])
    for i, m in enumerate(MOLDS, 2):
        for c, v in enumerate(m, 1):
            cell = inp.cell(row=i, column=c, value=v)
            cell.font = INPUT
            if c >= 4:
                cell.number_format = "mm/dd/yyyy"
    inp["F1"].comment = Comment("Press/trial slot from the press plan. Leave empty and the template plans TH1 as soon as unpack and gauge are done.", "PLM")
    style_rows(inp, 2, len(MOLDS) + 1, 6)

    hol = wb.create_sheet("Holidays")
    header(hol, 1, ["Plant holiday", "Name"], [16, 30])
    hol["C1"] = "Enter plant holidays here; all plan dates skip them."
    hol["C1"].font = Font(name=FONT, italic=True)
    for rr in range(2, 61):
        hol.cell(row=rr, column=1).number_format = "mm/dd/yyyy"

    pl = wb.create_sheet("Plan")
    cols = ["Mold", "Step ID", "Task", "Role", "Person", "Duration (wd)", "Start", "Finish", "PM plan start", "PM plan finish", "Delta finish (days)", "Checklist"]
    header(pl, 1, cols, [8, 8, 52, 22, 20, 11, 12, 12, 13, 13, 12, 50])
    r = 2
    steps = BLOCKS["LAUNCH"][1]
    for mi, m in enumerate(MOLDS):
        ir = mi + 2
        pos = {sid: r + 1 + k for k, (sid, *_) in enumerate(steps)}
        head = pl.cell(row=r, column=1, value=f"=Inputs!A{ir}")
        head.font = Font(name=FONT, bold=True, color="008000")
        pl.cell(row=r, column=3, value=f'=Inputs!B{ir}&": Mold "&Inputs!A{ir}&" ("&Inputs!C{ir}&")"').font = Font(name=FONT, bold=True)
        pl.cell(row=r, column=7, value=f"=MIN(G{r + 1}:G{r + len(steps)})")
        pl.cell(row=r, column=8, value=f"=MAX(H{r + 1}:H{r + len(steps)})")
        for col in range(1, 13):
            pl.cell(row=r, column=col).fill = BAND
        r += 1
        for sid, name, role, dur, rule, chk in steps:
            b = brow[sid]
            G = lambda s: f"G{pos[s]}"
            H = lambda s: f"H{pos[s]}"
            after = lambda ids: ",".join(f"WORKDAY({H(d)},1,{HOL})" for d in ids)
            k = rule[0]
            if k == "input":
                start = f"=Inputs!{ROLE_COL[rule[1]]}{ir}"
            elif k == "before":
                start = f"=WORKDAY({G(rule[1])},-F{r},{HOL})"
            elif k == "same":
                start = f"={G(rule[1])}"
            elif k == "after":
                start = f"=MAX({after(rule[1])})"
            elif k == "input_or_after":
                inc = f"Inputs!{ROLE_COL[rule[1]]}{ir}"
                start = f'=IF({inc}="",MAX({after(rule[2])}),{inc})'
            elif k == "max_input_after":
                start = f"=MAX(Inputs!{ROLE_COL[rule[1]]}{ir},{after(rule[2])})"
            else:
                start = "=MAX(" + ",".join(H(d) for d in rule[1]) + ")"
            vals = [f"=Inputs!A{ir}", f"=Blocks!B{b}", f"=Blocks!C{b}", f"=Blocks!D{b}",
                    f'=IFERROR(INDEX(Team!$B$2:$B$60,MATCH(D{r},Team!$A$2:$A$60,0)),"(assign)")',
                    f"=Blocks!E{b}", start, f"=IF(F{r}=0,G{r},WORKDAY(G{r},F{r}-1,{HOL}))"]
            for col, v in enumerate(vals, 1):
                cell = pl.cell(row=r, column=col, value=v)
                if col in (1, 2, 3, 4, 6):
                    cell.font = LINK
            pm = PM.get(m[0], {}).get(sid)
            if pm:
                for col, v in ((9, pm[0]), (10, pm[1])):
                    mm, dd = map(int, v.split("-"))
                    pl.cell(row=r, column=col, value=date(2026, mm, dd)).font = INPUT
                pl.cell(row=r, column=11, value=f'=IF(J{r}="","",H{r}-J{r})')
            pl.cell(row=r, column=12, value=f"=IF(Blocks!H{b}=\"\",\"\",Blocks!H{b})").font = LINK
            r += 1
    for rr in range(2, r):
        for col in (7, 8, 9, 10):
            pl.cell(row=rr, column=col).number_format = "mm/dd/yyyy"
        pl.cell(row=rr, column=11).number_format = '+0;-0;0'
    style_rows(pl, 2, r - 1, 12)
    for rr in range(2, r):
        if pl.cell(row=rr, column=2).value is None:
            for col in range(1, 13):
                pl.cell(row=rr, column=col).fill = BAND
                pl.cell(row=rr, column=col).font = Font(name=FONT, bold=True, color="008000" if col == 1 else None)
    pl.auto_filter.ref = f"A1:L{r - 1}"

    fd = wb.create_sheet("Findings")
    header(fd, 1, ["Area", "Found in the current .mpp", "How the template prevents it"], [14, 70, 70])
    for i, row in enumerate(FINDINGS, 2):
        for c, v in enumerate(row, 1):
            fd.cell(row=i, column=c, value=v)
    style_rows(fd, 2, len(FINDINGS) + 1, 3)
    for rr in range(2, len(FINDINGS) + 2):
        for c in (2, 3):
            fd.cell(row=rr, column=c).alignment = Alignment(wrap_text=True, vertical="top")

    wb.move_sheet("Plan", offset=-3)
    wb.save(XLSX)


# --- MS Project XML --------------------------------------------------------
def build_xml():
    import jpype
    import mpxj  # noqa: F401  (adds the jars)
    jh = os.environ["JAVA_HOME"]
    jpype.startJVM(jvmpath=os.path.join(jh, "lib", "server", "libjvm.so"))
    from java.time import LocalDateTime
    from org.mpxj import ConstraintType, Duration, ProjectFile, Relation, RelationType, TimeUnit
    from org.mpxj.mspdi import MSPDIWriter

    def ldt(d, h):
        return LocalDateTime.of(d.year, d.month, d.day, h, 0)

    pf = ProjectFile()
    pf.addDefaultBaseCalendar()
    props = pf.getProjectProperties()
    props.setName("G67 Timeline (from template)")
    props.setStartDate(ldt(min(m[3] for m in MOLDS) - timedelta(days=30), 8))
    res = {}
    for role, person, dept in TEAM:
        rs = pf.addResource()
        rs.setName(person)
        rs.setGroup(role)
        res[role] = rs

    top = pf.addTask()
    top.setName("G67 mold launches (block LAUNCH v1)")
    steps = BLOCKS["LAUNCH"][1]
    for mold, press, parts, deliv, gauge, th1 in MOLDS:
        inputs = {"delivery": deliv, "gauge": gauge, "th1": th1}
        s, f = schedule(steps, inputs)
        summ = top.addTask()
        summ.setName(f"{press}: Mold {mold} ({parts})")
        t = {}
        for sid, name, role, dur, rule, chk in steps:
            task = summ.addTask()
            task.setName(f"Mold {mold}: {name}")
            task.setText(1, sid)
            task.setText(2, role)
            task.setDuration(Duration.getInstance(dur, TimeUnit.DAYS))
            task.setStart(ldt(s[sid], 8))
            task.setFinish(ldt(f[sid], 8 if dur == 0 else 17))
            task.setMilestone(dur == 0)
            if chk:
                task.setNotes(chk)
            if dur:
                task.addResourceAssignment(res[role])
            k = rule[0]
            snet = {"input": inputs.get(rule[1]) if k == "input" else None,
                    "before": s[sid] if k == "before" else None,
                    "input_or_after": inputs.get(rule[1]) if k == "input_or_after" else None,
                    "max_input_after": inputs.get(rule[1]) if k == "max_input_after" else None}.get(k)
            if snet:
                task.setConstraintType(ConstraintType.START_NO_EARLIER_THAN)
                task.setConstraintDate(ldt(snet, 8))
            t[sid] = task
        links = []
        for sid, _, _, _, rule, _ in steps:
            k = rule[0]
            if k in ("after", "end"):
                links += [(d, sid, RelationType.FINISH_START) for d in rule[1]]
            elif k == "same":
                links.append((rule[1], sid, RelationType.START_START))
            elif k == "before":
                links.append((sid, rule[1], RelationType.FINISH_START))
            elif k == "max_input_after":
                links += [(d, sid, RelationType.FINISH_START) for d in rule[2]]
            elif k == "input_or_after" and not inputs.get(rule[1]):
                links += [(d, sid, RelationType.FINISH_START) for d in rule[2]]
        for a, b, typ in links:
            t[b].addPredecessor(Relation.Builder().predecessorTask(t[a]).type(typ).lag(Duration.getInstance(0, TimeUnit.DAYS)))
    pf.updateStructure()
    MSPDIWriter().write(pf, XML)


if __name__ == "__main__":
    build_xlsx()
    if os.environ.get("JAVA_HOME"):
        build_xml()
