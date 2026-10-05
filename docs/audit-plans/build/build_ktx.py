"""Builds TOC-PLM-00..05 as internal working documents (not document-controlled).

Styles and bullets come from the KTX template FM-QUA-0039-07; its controlled header
(number, revision), footer, approval and revision history are replaced by a plain
internal header/footer. The body is the content in content.py.
"""
import copy
import re
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

import content as C

TEMPLATE = "tpl/FM-QUA-0039-07.docx"
OUT = "../"
BULLET_NUM_ID = "52"
BODY_PT = 11
TABLE_PT = 9


def runs(par, text, size=None, bold=False, color=None):
    for part in re.split(r"(\*\*[^*]+\*\*)", text):
        if not part:
            continue
        b = part.startswith("**")
        r = par.add_run(part[2:-2] if b else part)
        r.bold = bold or b
        if size:
            r.font.size = Pt(size)
        if color:
            r.font.color.rgb = RGBColor.from_string(color)
    return par


def spacing(par, before=0, after=4, keep=False):
    pf = par.paragraph_format
    pf.space_before, pf.space_after = Pt(before), Pt(after)
    if keep:
        pf.keep_with_next = True


def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def row_flags(row, header=False):
    trPr = row._tr.get_or_add_trPr()
    cs = OxmlElement("w:cantSplit")
    trPr.append(cs)
    if header:
        th = OxmlElement("w:tblHeader")
        trPr.append(th)


def table(doc, widths, header, rows):
    t = doc.add_table(rows=0, cols=len(widths))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    all_rows = ([header] if header else []) + rows
    for ri, data in enumerate(all_rows):
        row = t.add_row()
        row_flags(row, header=bool(header) and ri == 0)
        for ci, val in enumerate(data):
            cell = row.cells[ci]
            cell.width = Inches(widths[ci])
            p = cell.paragraphs[0]
            spacing(p, 1, 1)
            runs(p, str(val), size=TABLE_PT, bold=bool(header) and ri == 0)
            if header and ri == 0:
                shade(cell, "D9D9D9")
    # fixed grid widths
    grid = t._tbl.tblGrid
    for gc, w in zip(grid.findall(qn("w:gridCol")), widths):
        gc.set(qn("w:w"), str(int(w * 1440)))
    doc.add_paragraph()
    return t


def bullet(doc, text):
    p = doc.add_paragraph(style="List Paragraph")
    pPr = p._p.get_or_add_pPr()
    numPr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl"); ilvl.set(qn("w:val"), "0")
    numId = OxmlElement("w:numId"); numId.set(qn("w:val"), BULLET_NUM_ID)
    numPr.append(ilvl); numPr.append(numId)
    pPr.insert(0, numPr) if pPr.find(qn("w:pStyle")) is None else pPr.find(qn("w:pStyle")).addnext(numPr)
    spacing(p, 0, 2)
    runs(p, text, size=BODY_PT)


def field(par, instr, size):
    """PAGE / NUMPAGES field, updated by Word on conversion."""
    for kind, text in (("begin", None), ("instr", instr), ("separate", None), ("text", "1"), ("end", None)):
        r = par.add_run()
        r.font.size = Pt(size)
        if kind == "instr":
            it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = f" {instr} "
            r._r.append(it)
        elif kind == "text":
            r.text = text
        else:
            fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), kind)
            r._r.append(fc)


def clear(part):
    for child in list(part._element):
        part._element.remove(child)


def fill_header(doc, d):
    """Internal working document: plain header and footer, no document-control fields."""
    sec = doc.sections[0]
    hdr, ftr = sec.header, sec.footer
    clear(hdr); clear(ftr)
    p = hdr.add_paragraph(); spacing(p, 0, 0)
    runs(p, f"KTX Toccoa \u00b7 {C.DEPT} \u00b7 {C.AREA}", size=8, color="6B7280")
    p = hdr.add_paragraph(); spacing(p, 0, 0)
    runs(p, d["title"], size=10, bold=True)
    p = hdr.add_paragraph(); spacing(p, 0, 8)
    runs(p, f"{d['no']} \u00b7 status {d.get('date', C.DATE)} \u00b7 {C.INTERNAL}", size=8, color="6B7280")
    pPr = p._p.get_or_add_pPr(); bdr = OxmlElement("w:pBdr"); b = OxmlElement("w:bottom")
    for k, v in (("val", "single"), ("sz", "4"), ("space", "4"), ("color", "9CA3AF")):
        b.set(qn(f"w:{k}"), v)
    bdr.append(b); pPr.append(bdr)
    p = ftr.add_paragraph(); spacing(p, 0, 0)
    p.paragraph_format.tab_stops.add_tab_stop(Inches(6.5), alignment=2)
    runs(p, f"Internal working document \u2013 not part of document control\t", size=8, color="6B7280")
    runs(p, "Page ", size=8, color="6B7280"); field(p, "PAGE", 8)
    runs(p, " of ", size=8, color="6B7280"); field(p, "NUMPAGES", 8)


def build(d):
    doc = Document(TEMPLATE)
    fill_header(doc, d)
    body = doc.element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)

    def h(text, level):
        p = doc.add_paragraph()
        spacing(p, 10 if level == 1 else 6, 4, keep=True)
        runs(p, text, size=12 if level == 1 else BODY_PT, bold=True)

    # Purpose / Authority, as the template opens
    h("Purpose / Scope:", 1)
    p = doc.add_paragraph(); spacing(p); runs(p, d["purpose"], size=BODY_PT)
    h("Owner / involved:", 1)
    for a in d.get("auth", C.AUTH):
        bullet(doc, a)

    for blk in d["blocks"]:
        kind = blk[0]
        if kind == "h1":
            h(blk[1], 1)
        elif kind == "h2":
            h(blk[1], 2)
        elif kind == "p":
            p = doc.add_paragraph(); spacing(p, 0, 6); runs(p, blk[1], size=BODY_PT)
        elif kind == "b":
            bullet(doc, blk[1])
        elif kind == "img":
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; spacing(p, 2, 6)
            p.add_run().add_picture(blk[1], width=Inches(blk[3] if len(blk) > 3 else 6.5))
        elif kind == "table":
            table(doc, blk[1], blk[2], blk[3])
        elif kind == "pb":
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    out = f"{OUT}{d['file']}.docx"
    doc.core_properties.title = d["title"]
    doc.core_properties.author = "KTX Toccoa, Process Development"
    doc.save(out)
    return out


if __name__ == "__main__":
    only = sys.argv[1:]
    for d in C.DOCS:
        if not only or d["no"] in only:
            print(build(d))
