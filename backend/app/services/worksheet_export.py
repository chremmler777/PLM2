"""xlsx of the worksheet as the browser shows it: visible columns and rows in
order, numbers and dates typed, flagged cells in the engineering Excel's
colours, the identity columns and header frozen. Text is never a formula."""
from datetime import date, datetime
from io import BytesIO
from typing import Optional

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
FLAG_FILLS = {"open": "FFFFFF00", "confirmed": "FFC6EFCE", "rejected": "FFF8CBAD"}
HEADER_FILL = "FF305496"


def _number(v):
    if isinstance(v, bool):
        return str(v)
    if isinstance(v, (int, float)):
        return v
    try:
        f = float(str(v).strip().replace(",", "."))
    except ValueError:
        return str(v)
    return int(f) if f.is_integer() else f


def _date(v):
    try:
        return datetime.fromisoformat(str(v).strip()[:19])
    except ValueError:
        try:
            return datetime.combine(date.fromisoformat(str(v).strip()[:10]), datetime.min.time())
        except ValueError:
            return str(v)


def _clean(text: str) -> str:
    """Drop control characters openpyxl refuses (they would fail the save)."""
    return ILLEGAL_CHARACTERS_RE.sub("", text)


def _typed(kind: str, v):
    if isinstance(v, str):
        v = _clean(v)
    if v is None or v == "":
        return None
    if kind == "number":
        return _number(v)
    if kind == "date":
        return _date(v)
    return str(v)


def _literal(cell) -> None:
    """Text starting like a formula stays literal text, never a formula."""
    if isinstance(cell.value, str) and cell.value[:1] in ("=", "+", "-", "@"):
        cell.data_type = "s"


def build_xlsx(columns: list[dict], rows: list[list[dict]], frozen_columns: int,
                sheet_title: str = "Worksheet") -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = _clean(sheet_title)[:31] or "Worksheet"
    labels = [_clean(c["label"]) for c in columns]
    for j, label in enumerate(labels, start=1):
        cell = ws.cell(row=1, column=j, value=label)
        _literal(cell)
        cell.font = Font(bold=True, color="FFFFFFFF")
        cell.fill = PatternFill("solid", fgColor=HEADER_FILL)
    widths = [len(label) for label in labels]
    for i, cells in enumerate(rows, start=2):
        for j, (col, data) in enumerate(zip(columns, cells), start=1):
            value = _typed(col.get("type", "text"), data.get("value"))
            cell = ws.cell(row=i, column=j, value=value)
            _literal(cell)
            if isinstance(value, datetime):
                cell.number_format = "yyyy-mm-dd"
            flag: Optional[str] = data.get("flag")
            if flag in FLAG_FILLS:
                cell.fill = PatternFill("solid", fgColor=FLAG_FILLS[flag])
            n = int(data.get("comments") or 0)
            if n > 0:
                cell.comment = Comment(f"{n} comment{'s' if n != 1 else ''} in PLM", "PLM")
            widths[j - 1] = max(widths[j - 1], min(len(str(value)) if value is not None else 0, 60))
    for j, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(j)].width = max(8, w + 2)
    ws.freeze_panes = f"{get_column_letter(frozen_columns + 1)}2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{len(rows) + 1}"
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
